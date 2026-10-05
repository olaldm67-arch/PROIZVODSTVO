import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import openpyxl
from tksheet import Sheet
import threading
import datetime
import os
import tempfile
import webbrowser
import html as html_lib


# Столбцы, по которым нужно считать итог
SUM_COLUMNS = [
    "Кол-во", "мкв",
    "Общ с грн", "Аванс грн", "Оплата грн",
    "Оплата2", "Монтаж", "Долг",
]

# Столбцы, для которых доступен выбор из списка уникальных значений
VALUE_FILTER_COLUMNS = ["город", "проф"]

# Числовые столбцы (для выравнивания и форматирования)
NUMERIC_COLUMNS = [
    "Кол-во", "мкв", "Скидка",
    "Общ с грн", "Аванс грн", "Оплата грн",
    "Оплата2", "Монтаж", "Долг",
]

# Столбец с плавающей точкой (1 знак после запятой)
FLOAT_COLUMNS = ["мкв"]

# Столбец, до которого включительно выводится печать
LAST_PRINT_COLUMN = "комент"

_NUM_LOWER = [c.lower() for c in NUMERIC_COLUMNS]
_FLOAT_LOWER = [c.lower() for c in FLOAT_COLUMNS]


class ExcelEditorApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Excel Viewer & Editor (32-bit compatible)")
        self.root.geometry("1200x700")

        # ---------- Данные ----------
        self.headers = []
        self.original_data = []
        self.displayed_rows = []
        self.cancel_flag = False
        self.progress_window = None

        # ---------- Виджеты фильтров и сумм ----------
        self.filter_cells = []
        self.filter_entries = []
        self.sum_labels = []

        self.selected_filter_values = {}
        self._last_selected_count = -1

        # ---------- Панель инструментов ----------
        toolbar = tk.Frame(root)
        toolbar.pack(fill=tk.X, padx=5, pady=5)

        tk.Button(toolbar, text="Загрузить Excel",
                  command=self.load_excel).pack(side=tk.LEFT, padx=5)
        tk.Button(toolbar, text="Сохранить Excel",
                  command=self.save_excel).pack(side=tk.LEFT, padx=5)
        tk.Button(toolbar, text="Сбросить все фильтры",
                  command=self.clear_all_filters).pack(side=tk.LEFT, padx=5)

        tk.Button(toolbar, text="📋 Копировать всё видимое",
                  bg='#d4e8d4',
                  command=self.copy_all_visible).pack(side=tk.LEFT, padx=5)
        tk.Button(toolbar, text="📋 Копировать выделенное (Ctrl+C)",
                  bg='#d4e4f0',
                  command=self.copy_selected).pack(side=tk.LEFT, padx=5)
        tk.Button(toolbar, text="🖨 Печать с просмотром",
                  bg='#f0e0a0',
                  command=self.print_preview).pack(side=tk.LEFT, padx=5)

        # ---------- Строка фильтров над таблицей ----------
        filter_row = tk.Frame(root, height=32, bg='#e0e0e0')
        filter_row.pack(fill=tk.X, padx=5)
        filter_row.pack_propagate(False)

        self.filter_canvas = tk.Canvas(filter_row, height=30, bg='#e0e0e0',
                                       highlightthickness=0)
        self.filter_canvas.pack(fill=tk.BOTH, expand=True)

        # ---------- Таблица ----------
        sheet_container = tk.Frame(root)
        sheet_container.pack(fill=tk.BOTH, expand=True, padx=5, pady=(0, 0))

        self.sheet = Sheet(sheet_container)
        self.sheet.pack(fill=tk.BOTH, expand=True)

        try:
            self.sheet.enable_bindings((
                "single_select",
                "row_select",
                "arrowkeys",
                "copy", "cut", "paste", "delete",
                "edit_cell",
                "rc_select", "right_click_popup_menu",
            ))
        except Exception:
            try:
                self.sheet.enable_bindings(
                    "single_select", "row_select", "arrowkeys",
                    "copy", "paste", "delete",
                    "edit_cell", "rc_select",
                    "right_click_popup_menu",
                )
            except Exception:
                try:
                    self.sheet.enable_bindings("all")
                except Exception as e:
                    print(f"enable_bindings error: {e}")

        try:
            self.sheet.extra_bindings("end_edit_cell", self.on_cell_edit)
        except Exception:
            pass

        self.root.bind_all("<Control-c>", self._on_ctrl_c)
        self.root.bind_all("<Control-C>", self._on_ctrl_c)
        self.root.bind_all("<Control-p>", self._on_ctrl_p)
        self.root.bind_all("<Control-P>", self._on_ctrl_p)

        try:
            self.sheet.bind("<Button-3>", self._on_sheet_right_click)
        except Exception:
            pass

        # ---------- Строка итогов под таблицей ----------
        sum_row = tk.Frame(root, height=28, bg='#f0f0f0')
        sum_row.pack(fill=tk.X, padx=5, pady=(0, 0))
        sum_row.pack_propagate(False)

        self.sum_canvas = tk.Canvas(sum_row, height=26, bg='#f0f0f0',
                                    highlightthickness=0)
        self.sum_canvas.pack(fill=tk.BOTH, expand=True)

        # ---------- Строка статуса ----------
        status_row = tk.Frame(root, height=24, bg='#333333')
        status_row.pack(fill=tk.X, padx=5, pady=(0, 5))
        status_row.pack_propagate(False)

        self.status_label = tk.Label(
            status_row,
            text="Выделено: 0 строк  •  Клик по номеру слева — вся строка, "
                 "протяжка или Shift+клик — несколько",
            font=("Arial", 9), bg='#333333', fg='#e0e0e0', anchor="w",
            padx=8)
        self.status_label.pack(fill=tk.BOTH, expand=True)

        self.root.after(100, self._sync_overlays)
        self.root.after(200, self._update_selection_status)

    # ====================================================================
    #                     ПЕЧАТЬ С ПРЕДВАРИТЕЛЬНЫМ ПРОСМОТРОМ
    # ====================================================================
    def _on_ctrl_p(self, event=None):
        try:
            widget = self.root.focus_get()
            if isinstance(widget, tk.Entry):
                return
        except Exception:
            pass
        self.print_preview()
        return "break"

    def _get_last_print_col_idx(self):
        """Индекс последнего столбца для печати (включительно) — «комент».
        Если такого столбца нет — печатаем все столбцы."""
        for i, h in enumerate(self.headers):
            if str(h).strip().lower() == LAST_PRINT_COLUMN.lower():
                return i
        return len(self.headers) - 1

    def _generate_print_html(self):
        """Формирует HTML только с таблицей — до столбца «комент» включительно.
        Без шапки, без инструкций, без мета-информации."""
        if not self.headers:
            return "<html><body></body></html>"

        last_col = self._get_last_print_col_idx()
        print_headers = self.headers[:last_col + 1]
        n_print_cols = len(print_headers)

        # Какие из печатаемых столбцов числовые
        num_cols = set()
        for i in range(n_print_cols):
            if str(self.headers[i]).strip().lower() in _NUM_LOWER:
                num_cols.add(i)

        css = """
        <style>
            @page { size: A4 landscape; margin: 8mm; }
            * { box-sizing: border-box; }
            html, body {
                font-family: "Segoe UI", Arial, sans-serif;
                font-size: 10px;
                margin: 0; padding: 0;
            }
            table {
                border-collapse: collapse;
                width: 100%;
                table-layout: auto;
            }
            th, td {
                border: 1px solid #999;
                padding: 3px 5px;
                vertical-align: top;
                word-wrap: break-word;
                max-width: 240px;
            }
            th {
                background: #4a6fa5;
                color: white;
                font-weight: bold;
                text-align: center;
                font-size: 10px;
            }
            tr:nth-child(even) td { background: #f5f8fc; }
            td.num { text-align: right; white-space: nowrap; }
            td.txt { text-align: left; }
            tfoot td {
                background: #ffe9a8;
                font-weight: bold;
                border-top: 2px solid #444;
                text-align: right;
                white-space: nowrap;
            }
            tfoot td.lbl { text-align: left; color: #333; }
        </style>
        """

        parts = []
        parts.append('<!DOCTYPE html>')
        parts.append('<html lang="ru"><head><meta charset="utf-8">')
        parts.append('<title>Таблица</title>')
        parts.append(css)
        parts.append('</head><body>')

        parts.append('<table>')

        # Заголовки
        parts.append('<thead><tr>')
        for h in print_headers:
            parts.append(f'<th>{html_lib.escape(str(h))}</th>')
        parts.append('</tr></thead>')

        # Данные
        parts.append('<tbody>')
        for row in self.displayed_rows:
            parts.append('<tr>')
            for i in range(n_print_cols):
                v = row[i] if i < len(row) else ""
                v_safe = html_lib.escape(str(v))
                cls = "num" if i in num_cols else "txt"
                parts.append(f'<td class="{cls}">{v_safe}</td>')
            parts.append('</tr>')
        parts.append('</tbody>')

        # Итоги — только для столбцов до last_col включительно
        totals_text = self._compute_totals_for_print(last_col)
        if any(totals_text.get(i, "") for i in range(n_print_cols)):
            parts.append('<tfoot><tr>')
            for i in range(n_print_cols):
                txt = totals_text.get(i, "")
                if txt:
                    parts.append(f'<td class="num">{html_lib.escape(txt)}</td>')
                else:
                    parts.append('<td class="lbl"></td>')
            parts.append('</tr></tfoot>')

        parts.append('</table>')
        parts.append('</body></html>')

        return "\n".join(parts)

    def _compute_totals_for_print(self, last_col):
        """Итоги только по столбцам до last_col включительно."""
        result = {}
        if not self.headers:
            return result

        for col_idx in range(last_col + 1):
            if col_idx >= len(self.headers):
                break
            header = self.headers[col_idx]
            if header not in SUM_COLUMNS:
                continue

            total = 0.0
            has = False
            for row in self.displayed_rows:
                if col_idx >= len(row):
                    continue
                s = str(row[col_idx]).strip()
                if not s:
                    continue
                try:
                    total += float(s.replace(" ", "").replace(",", "."))
                    has = True
                except (ValueError, AttributeError):
                    pass

            if not has:
                continue

            if self._is_float_column(col_idx):
                result[col_idx] = f"{total:.1f}"
            else:
                result[col_idx] = f"{int(round(total)):,}".replace(",", " ")
        return result

    def print_preview(self):
        """Формирует HTML и открывает его в браузере по умолчанию."""
        if not self.headers:
            messagebox.showwarning(
                "Внимание",
                "Сначала загрузите файл Excel."
            )
            return

        if not self.displayed_rows:
            if not messagebox.askyesno(
                    "Нет данных",
                    "Текущий фильтр не нашёл ни одной строки.\n"
                    "Распечатать пустую таблицу с заголовками?"):
                return

        try:
            html_doc = self._generate_print_html()

            tmp_dir = tempfile.gettempdir()
            ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"table_print_{ts}.html"
            filepath = os.path.join(tmp_dir, filename)

            with open(filepath, "w", encoding="utf-8") as f:
                f.write(html_doc)

            webbrowser.open(f"file:///{filepath.replace(os.sep, '/')}")

        except Exception as e:
            messagebox.showerror("Ошибка",
                                 f"Не удалось подготовить печать:\n{e}")

    # ====================================================================
    #              Хелперы для числовых столбцов
    # ====================================================================
    def _is_numeric_column(self, col_idx):
        if col_idx >= len(self.headers):
            return False
        return str(self.headers[col_idx]).strip().lower() in _NUM_LOWER

    def _is_float_column(self, col_idx):
        if col_idx >= len(self.headers):
            return False
        return str(self.headers[col_idx]).strip().lower() in _FLOAT_LOWER

    def _format_number_value(self, value, col_idx):
        try:
            s = str(value).strip().replace(" ", "").replace(",", ".")
            num = float(s)
        except (ValueError, AttributeError):
            return str(value) if value is not None else ""

        if self._is_float_column(col_idx):
            return f"{num:.1f}"
        else:
            return str(int(round(num)))

    def _is_value_filter_column(self, col_idx):
        if col_idx >= len(self.headers):
            return False
        header = str(self.headers[col_idx]).strip().lower()
        return header in [v.lower() for v in VALUE_FILTER_COLUMNS]

    # ====================================================================
    #              Обновление статуса выделенных строк
    # ====================================================================
    def _update_selection_status(self):
        try:
            selected = self._get_selected_row_indices()
            count = len(selected)
            if count != self._last_selected_count:
                self._last_selected_count = count
                if count == 0:
                    self.status_label.config(
                        text="Выделено: 0 строк  •  Клик по номеру слева — вся строка, "
                             "протяжка или Shift+клик — несколько",
                        fg='#e0e0e0')
                else:
                    self.status_label.config(
                        text=f"✓ Выделено строк: {count}  •  "
                             f"Нажмите Ctrl+C или «Копировать выделенное»",
                        fg='#7fdc7f')
        except Exception as e:
            print(f"Selection status error: {e}")

        self.root.after(200, self._update_selection_status)

    def _on_ctrl_c(self, event=None):
        try:
            widget = self.root.focus_get()
            if isinstance(widget, tk.Entry):
                return
        except Exception:
            pass
        self.copy_selected()
        return "break"

    def _on_sheet_right_click(self, event):
        menu = tk.Menu(self.root, tearoff=0)
        menu.add_command(label="📋 Копировать выделенные строки",
                         command=self.copy_selected)
        menu.add_command(label="📋 Копировать всё видимое",
                         command=self.copy_all_visible)
        menu.add_separator()
        menu.add_command(label="🖨 Печать с предварительным просмотром",
                         command=self.print_preview)
        menu.add_separator()
        menu.add_command(label="Снять выделение",
                         command=self._deselect_all)
        try:
            menu.tk_popup(event.x_root, event.y_root)
        finally:
            menu.grab_release()

    def _deselect_all(self):
        try:
            self.sheet.deselect("all")
        except Exception:
            try:
                self.sheet.deselect()
            except Exception:
                pass

    # ====================================================================
    #           Получение выделенных строк
    # ====================================================================
    def _get_selected_row_indices(self):
        row_indices = set()

        try:
            rows = self.sheet.get_selected_rows()
            if rows:
                for r in rows:
                    try:
                        row_indices.add(int(r))
                    except (ValueError, TypeError):
                        pass
        except Exception as e:
            print(f"get_selected_rows failed: {e}")

        if not row_indices:
            try:
                sel = self.sheet.get_currently_selected()
                if sel is not None:
                    rows = getattr(sel, "rows", None)
                    if rows:
                        for r in rows:
                            try:
                                row_indices.add(int(r))
                            except (ValueError, TypeError):
                                pass
            except Exception as e:
                print(f"get_currently_selected failed: {e}")

        if not row_indices:
            try:
                cells = self.sheet.get_selected_cells()
                if cells:
                    for c in cells:
                        if isinstance(c, (list, tuple)) and len(c) >= 1:
                            try:
                                row_indices.add(int(c[0]))
                            except (ValueError, TypeError):
                                pass
            except Exception as e:
                print(f"get_selected_cells failed: {e}")

        return row_indices

    # ====================================================================
    #                       Копирование в буфер обмена
    # ====================================================================
    def _rows_to_clipboard_text(self, rows, with_headers=True):
        lines = []
        if with_headers and self.headers:
            lines.append("\t".join(self.headers))
        for row in rows:
            safe = [str(cell).replace("\n", " ").replace("\r", " ")
                    for cell in row]
            lines.append("\t".join(safe))
        return "\n".join(lines)

    def _push_to_clipboard(self, text):
        try:
            self.root.clipboard_clear()
            self.root.clipboard_append(text)
            self.root.update()
            return True
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось скопировать:\n{e}")
            return False

    def copy_all_visible(self):
        if not self.displayed_rows:
            messagebox.showwarning("Внимание", "Нет данных для копирования")
            return

        text = self._rows_to_clipboard_text(self.displayed_rows, with_headers=True)
        if self._push_to_clipboard(text):
            messagebox.showinfo(
                "Скопировано",
                f"Скопировано {len(self.displayed_rows)} строк с заголовками.\n"
                f"Вставьте их в Excel через Ctrl+V."
            )

    def copy_selected(self):
        if not self.displayed_rows:
            messagebox.showwarning("Внимание", "Нет данных для копирования")
            return

        data_rows = self._get_selected_row_indices()

        if not data_rows:
            messagebox.showwarning(
                "Строки не выделены",
                "Как выделить строки:\n\n"
                "1. Нажмите ЛКМ на СЕРЫЙ НОМЕР строки слева от таблицы.\n\n"
                "2. Зажмите ЛКМ и протяните вниз для нескольких строк.\n\n"
                "3. Или кликните по первому номеру, затем Shift+клик по последнему.\n\n"
                "4. Или зажмите Ctrl и кликайте по отдельным номерам.\n\n"
                "Чтобы скопировать всё — кнопка «📋 Копировать всё видимое»."
            )
            return

        selected_data = []
        for r in sorted(data_rows):
            if 0 <= r < len(self.displayed_rows):
                selected_data.append(self.displayed_rows[r])

        if not selected_data:
            messagebox.showwarning("Внимание",
                                   "В выделении нет строк с данными.")
            return

        text = self._rows_to_clipboard_text(selected_data, with_headers=True)
        if self._push_to_clipboard(text):
            messagebox.showinfo(
                "Скопировано",
                f"Скопировано {len(selected_data)} выделенных строк."
            )

    # ====================================================================
    #           Синхронизация строки фильтров и строки сумм
    # ====================================================================
    def _sync_overlays(self):
        if not self.headers or not self.filter_cells:
            self.root.after(150, self._sync_overlays)
            return

        try:
            try:
                rh_width = self.sheet.row_index_width()
            except Exception:
                rh_width = 40

            try:
                canvas = self.sheet.MT.canvas
                x_offset = canvas.canvasx(0)
            except Exception:
                x_offset = 0

            filter_w = self.filter_canvas.winfo_width()
            if filter_w <= 1:
                filter_w = self.root.winfo_width() - 20
            sum_w = self.sum_canvas.winfo_width()
            if sum_w <= 1:
                sum_w = self.root.winfo_width() - 20

            MIN_FILTER_WIDTH = 46
            MIN_SUM_WIDTH = 30

            x = rh_width - x_offset
            for i in range(len(self.headers)):
                try:
                    w = self.sheet.column_width(i)
                except Exception:
                    w = 100

                cell = self.filter_cells[i]
                vis_left = max(0, x)
                vis_right = min(x + w, filter_w)
                visible_w = vis_right - vis_left
                if visible_w >= MIN_FILTER_WIDTH:
                    cell.place(x=vis_left, y=2, width=visible_w, height=24)
                else:
                    cell.place_forget()

                if i < len(self.sum_labels):
                    lbl = self.sum_labels[i]
                    vis_left2 = max(0, x)
                    vis_right2 = min(x + w, sum_w)
                    visible_w2 = vis_right2 - vis_left2
                    if visible_w2 >= MIN_SUM_WIDTH:
                        lbl.place(x=vis_left2, y=2, width=visible_w2, height=22)
                    else:
                        lbl.place_forget()

                x += w

        except Exception as e:
            print(f"Sync error: {e}")

        self.root.after(50, self._sync_overlays)

    def _build_filter_cells(self):
        for cell in self.filter_cells:
            try:
                cell.destroy()
            except Exception:
                pass
        self.filter_cells.clear()
        self.filter_entries.clear()
        self.selected_filter_values.clear()

        for i, header in enumerate(self.headers):
            cell = tk.Frame(self.filter_canvas, bd=1, relief=tk.SOLID, bg='white')
            cell.columnconfigure(0, weight=1, minsize=15)
            cell.columnconfigure(1, weight=0, minsize=22)
            cell.rowconfigure(0, weight=1)

            entry = tk.Entry(cell, font=("Arial", 8), bd=0, bg='white',
                             highlightthickness=0)
            entry.grid(row=0, column=0, sticky="nsew", padx=(2, 0), pady=1)
            entry.bind("<KeyRelease>", lambda e: self.apply_all_filters())

            btn = tk.Button(cell, text="▼", font=("Arial", 9, "bold"),
                            bd=0, relief=tk.FLAT, bg='#c8c8c8',
                            activebackground='#909090', cursor="hand2",
                            command=lambda idx=i: self._show_column_menu(idx))
            btn.grid(row=0, column=1, sticky="nsew")

            self.filter_cells.append(cell)
            self.filter_entries.append(entry)

    def _build_sum_labels(self):
        for lbl in self.sum_labels:
            try:
                lbl.destroy()
            except Exception:
                pass
        self.sum_labels.clear()

        for i, header in enumerate(self.headers):
            lbl = tk.Label(self.sum_canvas, text="", font=("Arial", 8, "bold"),
                           bg='#f0f0f0', fg='blue', anchor="e",
                           bd=0, relief=tk.FLAT, padx=4)
            self.sum_labels.append(lbl)

    def _show_column_menu(self, col_idx):
        menu = tk.Menu(self.root, tearoff=0)
        header_name = self.headers[col_idx] if col_idx < len(self.headers) else f"Столбец {col_idx}"

        menu.add_command(label=f"Столбец: {header_name}", state="disabled")
        menu.add_separator()
        menu.add_command(label="Сортировать ↑ (A→Я, 0→9)",
                         command=lambda: self.sort_column(col_idx, ascending=True))
        menu.add_command(label="Сортировать ↓ (Я→A, 9→0)",
                         command=lambda: self.sort_column(col_idx, ascending=False))
        menu.add_separator()

        if self._is_value_filter_column(col_idx):
            active = col_idx in self.selected_filter_values and self.selected_filter_values[col_idx]
            label = "📋 Выбрать из списка значений" + ("  ✓" if active else "")
            menu.add_command(label=label,
                             command=lambda: self._open_value_selector(col_idx))
            menu.add_separator()

        menu.add_command(label="Сбросить фильтр этого столбца",
                         command=lambda: self.clear_column_filter(col_idx))

        try:
            btn = self.filter_cells[col_idx].winfo_children()[-1]
            x = btn.winfo_rootx()
            y = btn.winfo_rooty() + btn.winfo_height()
            menu.tk_popup(x, y)
        except Exception:
            menu.tk_popup(self.root.winfo_pointerx(), self.root.winfo_pointery())

    # ====================================================================
    #           Окно выбора уникальных значений столбца
    # ====================================================================
    def _open_value_selector(self, col_idx):
        header = self.headers[col_idx]

        unique_set = set()
        for row in self.original_data:
            if col_idx < len(row):
                v = str(row[col_idx]).strip()
                if v:
                    unique_set.add(v)
        unique_sorted = sorted(unique_set, key=lambda s: s.lower())

        if not unique_sorted:
            messagebox.showinfo("Пусто",
                                f"В столбце «{header}» нет значений для выбора.")
            return

        win = tk.Toplevel(self.root)
        win.title(f"Выбор значений: {header}")
        win.geometry("360x560")
        win.transient(self.root)
        win.grab_set()

        try:
            cell = self.filter_cells[col_idx]
            x = cell.winfo_rootx()
            y = cell.winfo_rooty() + cell.winfo_height()
            win.geometry(f"+{x}+{y}")
        except Exception:
            pass

        tk.Label(win,
                 text=f"Столбец «{header}» — {len(unique_sorted)} значений",
                 font=("Arial", 10, "bold")).pack(pady=(8, 4))

        search_frame = tk.Frame(win)
        search_frame.pack(fill=tk.X, padx=8, pady=(0, 4))

        tk.Label(search_frame, text="Поиск:").pack(side=tk.LEFT)
        search_var = tk.StringVar()
        search_entry = tk.Entry(search_frame, textvariable=search_var)
        search_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(4, 0))
        search_entry.focus_set()

        list_frame = tk.Frame(win)
        list_frame.pack(fill=tk.BOTH, expand=True, padx=8, pady=4)

        sb = tk.Scrollbar(list_frame)
        sb.pack(side=tk.RIGHT, fill=tk.Y)

        listbox = tk.Listbox(list_frame, selectmode=tk.EXTENDED,
                             yscrollcommand=sb.set, font=("Arial", 10),
                             activestyle="none")
        listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        sb.config(command=listbox.yview)

        current_set = self.selected_filter_values.get(col_idx, set())

        def fill_listbox(filter_text=""):
            listbox.delete(0, tk.END)
            ft = filter_text.strip().lower()
            for v in unique_sorted:
                if not ft or ft in v.lower():
                    listbox.insert(tk.END, v)
            if current_set:
                items = listbox.get(0, tk.END)
                for i, item in enumerate(items):
                    if item in current_set:
                        listbox.selection_set(i)

        fill_listbox()

        def on_search_change(*args):
            fill_listbox(search_var.get())

        search_var.trace("w", on_search_change)

        btn_frame = tk.Frame(win)
        btn_frame.pack(fill=tk.X, padx=8, pady=8)

        counter_label = tk.Label(win, text="", font=("Arial", 9), fg='#333')
        counter_label.pack(pady=(0, 4))

        def update_counter(*_):
            n = len(listbox.curselection())
            total = listbox.size()
            counter_label.config(text=f"Выбрано: {n} из {total}")

        listbox.bind("<<ListboxSelect>>", update_counter)
        update_counter()

        def select_all():
            listbox.selection_set(0, tk.END)
            update_counter()

        def clear_selection():
            listbox.selection_clear(0, tk.END)
            update_counter()

        def apply_selection():
            sel = listbox.curselection()
            chosen = {listbox.get(i) for i in sel}
            total_items = set(listbox.get(0, tk.END))

            if not chosen:
                self.selected_filter_values.pop(col_idx, None)
            elif chosen == total_items and len(total_items) == len(unique_set):
                self.selected_filter_values.pop(col_idx, None)
            else:
                self.selected_filter_values[col_idx] = set(chosen)

            self._refresh_filter_cell_highlight(col_idx)
            win.destroy()
            self.apply_all_filters()

        def reset_filter():
            self.selected_filter_values.pop(col_idx, None)
            self._refresh_filter_cell_highlight(col_idx)
            win.destroy()
            self.apply_all_filters()

        tk.Button(btn_frame, text="Выбрать всё", width=12,
                  command=select_all).pack(side=tk.LEFT, padx=2)
        tk.Button(btn_frame, text="Снять выделение", width=14,
                  command=clear_selection).pack(side=tk.LEFT, padx=2)
        tk.Button(btn_frame, text="Снять фильтр", width=12,
                  bg='#f0d0d0', command=reset_filter).pack(side=tk.LEFT, padx=2)

        bottom = tk.Frame(win)
        bottom.pack(fill=tk.X, padx=8, pady=(0, 8))

        tk.Button(bottom, text="Отмена", width=10,
                  command=win.destroy).pack(side=tk.RIGHT, padx=2)
        tk.Button(bottom, text="Применить", width=12, bg='#a8e8a8',
                  command=apply_selection).pack(side=tk.RIGHT, padx=2)

        win.bind("<Return>", lambda e: apply_selection())
        win.bind("<Escape>", lambda e: win.destroy())

    def _refresh_filter_cell_highlight(self, col_idx):
        if col_idx >= len(self.filter_entries):
            return
        try:
            entry = self.filter_entries[col_idx]
            active = (col_idx in self.selected_filter_values
                      and self.selected_filter_values[col_idx])
            entry.config(bg='#cfe8ff' if active else 'white')
        except Exception:
            pass

    # ====================================================================
    #                          Сортировка / фильтрация
    # ====================================================================
    def sort_column(self, col_idx, ascending=True):
        if not self.original_data:
            return

        def key_func(row):
            val = row[col_idx] if col_idx < len(row) else ""
            s = str(val).strip()
            if s == "":
                return (2, 0, "")
            try:
                num = float(s.replace(",", ".").replace(" ", ""))
                return (0, num, "")
            except (ValueError, AttributeError):
                pass
            return (1, 0, s.lower())

        self.original_data.sort(key=key_func, reverse=not ascending)
        self.apply_all_filters()

    def apply_all_filters(self, event=None):
        if not self.headers:
            return

        filtered = self.original_data

        for col_idx, entry in enumerate(self.filter_entries):
            text = entry.get().strip().lower()
            if text:
                filtered = [
                    row for row in filtered
                    if text in (row[col_idx] if col_idx < len(row) else "").lower()
                ]

            values = self.selected_filter_values.get(col_idx)
            if values:
                filtered = [
                    row for row in filtered
                    if (str(row[col_idx]).strip() if col_idx < len(row) else "") in values
                ]

        self._display_data(filtered)

    def clear_column_filter(self, col_idx):
        if 0 <= col_idx < len(self.filter_entries):
            self.filter_entries[col_idx].delete(0, tk.END)
        if col_idx in self.selected_filter_values:
            self.selected_filter_values.pop(col_idx, None)
            self._refresh_filter_cell_highlight(col_idx)
        self.apply_all_filters()

    def clear_all_filters(self):
        for entry in self.filter_entries:
            entry.delete(0, tk.END)
        self.selected_filter_values.clear()
        for i in range(len(self.filter_entries)):
            self._refresh_filter_cell_highlight(i)
        self.apply_all_filters()

    # ====================================================================
    #              Отображение данных
    # ====================================================================
    def _display_data(self, data):
        self.displayed_rows = data

        if data:
            display = data
        else:
            display = [[""] * len(self.headers)] if self.headers else []

        try:
            self.sheet.set_sheet_data(display,
                                      reset_col_positions=False,
                                      reset_row_positions=True)
            self.sheet.headers(self.headers)
        except Exception as e:
            print(f"Display error: {e}")

        self._apply_alignment()
        self._update_sums()

    def _apply_alignment(self):
        if not self.headers:
            return
        try:
            num_cols = [i for i in range(len(self.headers))
                        if self._is_numeric_column(i)]
            if not num_cols:
                return
            applied = False
            try:
                self.sheet.align_columns(columns=num_cols, align="e")
                applied = True
            except Exception:
                pass
            if not applied:
                try:
                    self.sheet.align_columns(num_cols, "e")
                    applied = True
                except Exception:
                    pass
            if not applied:
                try:
                    self.sheet.align(columns=num_cols, align="e")
                except Exception:
                    pass
        except Exception as e:
            print(f"align error: {e}")

    def _autofit_columns(self):
        if not self.headers:
            return
        try:
            widths = []
            for col_idx in range(len(self.headers)):
                max_len = len(str(self.headers[col_idx]))
                for row in self.displayed_rows[:1000]:
                    if col_idx < len(row):
                        v = str(row[col_idx])
                        if len(v) > max_len:
                            max_len = len(v)
                w = max(60, min(400, int(max_len * 7.5) + 18))
                widths.append(w)

            applied = False
            try:
                self.sheet.set_column_widths(widths)
                applied = True
            except Exception:
                pass
            if not applied:
                for i, w in enumerate(widths):
                    try:
                        self.sheet.column_width(i, w)
                    except Exception:
                        pass
        except Exception as e:
            print(f"autofit error: {e}")

    # ====================================================================
    #                          Подсчёт сумм
    # ====================================================================
    def _update_sums(self):
        if not self.headers or not self.sum_labels:
            return

        for lbl in self.sum_labels:
            lbl.config(text="")

        for col_idx, header in enumerate(self.headers):
            if header not in SUM_COLUMNS:
                continue
            if col_idx >= len(self.sum_labels):
                continue

            total = 0.0
            has_numeric = False
            non_empty = 0

            for row in self.displayed_rows:
                if col_idx >= len(row):
                    continue
                s = str(row[col_idx]).strip()
                if not s:
                    continue
                non_empty += 1
                try:
                    num = float(s.replace(" ", "").replace(",", "."))
                    total += num
                    has_numeric = True
                except (ValueError, AttributeError):
                    pass

            lbl = self.sum_labels[col_idx]
            if has_numeric:
                if self._is_float_column(col_idx):
                    text = f"{total:.1f}"
                else:
                    text = f"{int(round(total)):,}".replace(",", " ")
                lbl.config(text=text, fg='blue')
            elif non_empty > 0:
                lbl.config(text=f"[{non_empty}]", fg='#888888')

    # ====================================================================
    #                   Отслеживание правок ячеек
    # ====================================================================
    def on_cell_edit(self, event):
        try:
            row = event.row
            col = event.column
            if 0 <= row < len(self.displayed_rows):
                target_row = self.displayed_rows[row]
                val = self.sheet.get_cell_data(event.row, col)
                if self._is_numeric_column(col):
                    formatted = self._format_number_value(val, col)
                    target_row[col] = formatted
                    try:
                        self.sheet.set_cell_data(event.row, col, formatted)
                    except Exception:
                        pass
                else:
                    target_row[col] = str(val) if val is not None else ""
                self._update_sums()
        except Exception as e:
            print(f"Edit error: {e}")

    # ====================================================================
    #                          Загрузка файла
    # ====================================================================
    def load_excel(self):
        file_path = filedialog.askopenfilename(
            title="Выберите файл Excel",
            filetypes=[("Excel files", "*.xlsx *.xls")]
        )
        if not file_path:
            return

        self.cancel_flag = False
        self.show_progress_window()

        thread = threading.Thread(target=self._load_excel_thread,
                                  args=(file_path,), daemon=True)
        thread.start()

    def show_progress_window(self):
        self.progress_window = tk.Toplevel(self.root)
        self.progress_window.title("Загрузка файла...")
        self.progress_window.geometry("420x160")
        self.progress_window.transient(self.root)
        self.progress_window.grab_set()
        self.progress_window.resizable(False, False)

        self.progress_window.update_idletasks()
        x = self.root.winfo_x() + (self.root.winfo_width() - 420) // 2
        y = self.root.winfo_y() + (self.root.winfo_height() - 160) // 2
        self.progress_window.geometry(f"+{x}+{y}")

        self.progress_label = tk.Label(self.progress_window,
                                       text="Открытие файла...",
                                       font=("Arial", 10))
        self.progress_label.pack(pady=(15, 5))

        self.progress_bar = ttk.Progressbar(self.progress_window,
                                            length=380,
                                            mode='determinate',
                                            maximum=100)
        self.progress_bar.pack(pady=5)

        self.progress_percent = tk.Label(self.progress_window,
                                         text="0%", font=("Arial", 9))
        self.progress_percent.pack(pady=2)

        tk.Button(self.progress_window, text="Отмена",
                  command=self.cancel_load).pack(pady=5)
        self.progress_window.protocol("WM_DELETE_WINDOW", self.cancel_load)

    def cancel_load(self):
        self.cancel_flag = True

    def _load_excel_thread(self, file_path):
        try:
            self.root.after(0, lambda: self.progress_label.config(
                text="Чтение структуры файла..."))

            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            sheet = wb.active

            max_row = sheet.max_row or 0
            max_col = sheet.max_column or 0

            if max_row == 0 or max_col == 0:
                wb.close()
                self.root.after(0, lambda: self._finish_load([], []))
                return

            self.root.after(0, lambda: self.progress_label.config(
                text="Чтение заголовков..."))
            headers = []
            for row in sheet.iter_rows(min_row=1, max_row=1, values_only=True):
                for cell_value in row:
                    headers.append(str(cell_value) if cell_value is not None else "")
                break

            numeric_flags = [str(h).strip().lower() in _NUM_LOWER for h in headers]
            float_flags = [str(h).strip().lower() in _FLOAT_LOWER for h in headers]

            self.root.after(0, lambda: self.progress_label.config(
                text="Чтение данных..."))
            data = []
            total_data_rows = max_row - 1
            update_every = max(1, total_data_rows // 100)

            for row_idx, row in enumerate(
                    sheet.iter_rows(min_row=2, values_only=True), start=1):
                if self.cancel_flag:
                    wb.close()
                    self.root.after(0, self._cancel_finish)
                    return

                row_data = []
                for col_idx, v in enumerate(row):
                    if v is None:
                        row_data.append("")
                        continue
                    if isinstance(v, (datetime.datetime, datetime.date)):
                        row_data.append(v.strftime("%Y-%m-%d"))
                        continue
                    if col_idx < len(numeric_flags) and numeric_flags[col_idx]:
                        try:
                            num = float(str(v).replace(" ", "").replace(",", "."))
                            if col_idx < len(float_flags) and float_flags[col_idx]:
                                row_data.append(f"{num:.1f}")
                            else:
                                row_data.append(str(int(round(num))))
                        except (ValueError, AttributeError):
                            row_data.append(str(v))
                    else:
                        row_data.append(str(v))

                while len(row_data) < max_col:
                    row_data.append("")
                data.append(row_data)

                if row_idx % update_every == 0 or row_idx == total_data_rows:
                    percent = int((row_idx / total_data_rows) * 100) if total_data_rows > 0 else 100
                    self.root.after(0, lambda p=percent: self._update_progress(p))

            wb.close()
            self.root.after(0, lambda: self._finish_load(headers, data))

        except Exception as e:
            self.root.after(0, lambda err=e: self._load_error(err))

    def _update_progress(self, percent):
        if self.progress_window and self.progress_window.winfo_exists():
            self.progress_bar['value'] = percent
            self.progress_percent.config(text=f"{percent}%")

    def _finish_load(self, headers, data):
        if self.progress_window and self.progress_window.winfo_exists():
            self.progress_window.grab_release()
            self.progress_window.destroy()
        self.progress_window = None

        self.headers = headers
        self.original_data = data
        self.displayed_rows = data

        self._build_filter_cells()
        self._build_sum_labels()
        self._display_data(data)

        self.root.after(50, self._autofit_columns)
        self.root.after(150, self._sync_overlays)

        messagebox.showinfo("Успех", f"Загружено {len(data)} строк")

    def _cancel_finish(self):
        if self.progress_window and self.progress_window.winfo_exists():
            self.progress_window.grab_release()
            self.progress_window.destroy()
        self.progress_window = None

    def _load_error(self, e):
        if self.progress_window and self.progress_window.winfo_exists():
            self.progress_window.grab_release()
            self.progress_window.destroy()
        self.progress_window = None
        messagebox.showerror("Ошибка", f"Не удалось загрузить файл:\n{e}")

    # ====================================================================
    #                          Сохранение файла
    # ====================================================================
    def save_excel(self):
        if not self.headers:
            messagebox.showwarning("Внимание", "Сначала загрузите файл Excel")
            return

        file_path = filedialog.asksaveasfilename(
            title="Сохранить как",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")]
        )
        if not file_path:
            return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active

            for col_idx, header in enumerate(self.headers, start=1):
                ws.cell(row=1, column=col_idx, value=header)

            for row_idx, row_data in enumerate(self.original_data, start=2):
                for col_idx, value in enumerate(row_data, start=1):
                    if (col_idx - 1) < len(self.headers) and self._is_numeric_column(col_idx - 1):
                        try:
                            num_str = str(value).replace(" ", "").replace(",", ".")
                            if "." in num_str:
                                ws.cell(row=row_idx, column=col_idx,
                                        value=float(num_str))
                            else:
                                ws.cell(row=row_idx, column=col_idx,
                                        value=int(num_str))
                            continue
                        except (ValueError, AttributeError):
                            pass
                    ws.cell(row=row_idx, column=col_idx, value=value)

            wb.save(file_path)
            messagebox.showinfo("Успех", "Данные сохранены!")
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось сохранить файл:\n{e}")


if __name__ == "__main__":
    root = tk.Tk()
    app = ExcelEditorApp(root)
    root.mainloop()