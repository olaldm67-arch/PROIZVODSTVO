#!/usr/bin/env -S python3
# -*- coding: utf-8 -*-
"""
Утилита выгрузки структуры таблиц MySQL в Excel и PDF.

Создаёт два файла:
  db_structure.xlsx — лист "Обзор" + по листу на каждую таблицу
  db_structure.pdf  — единый документ со всеми таблицами
"""

import sys
import os
import pymysql
from datetime import datetime

# --- Excel ---
try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Не установлен openpyxl. Установите: pip install openpyxl")
    sys.exit(1)

# --- PDF ---
try:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                     Paragraph, Spacer, PageBreak)
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
except ImportError:
    print("Не установлен reportlab. Установите: pip install reportlab")
    sys.exit(1)


# ----------------------------------------------------------------------
# Параметры подключения к БД
# ----------------------------------------------------------------------
DB_HOST = '192.168.0.10'
DB_USER = 'workshop'
DB_PASSWORD = 'w0rK5h0p'
DB_NAME = 'workshop'
DB_CHARSET = 'utf8'

# Таблицы, структуру которых нужно выгрузить
TABLES = ['WORKER', 'WORKPLACE', 'WORKER_TIME', 'CCALC', 'DETAL_TIME']

# Имена выходных файлов
OUT_XLSX = 'db_structure.xlsx'
OUT_PDF = 'db_structure.pdf'


# ----------------------------------------------------------------------
# Получение структуры таблиц из information_schema
# ----------------------------------------------------------------------
def get_table_structure(cursor, table_name):
    """
    Возвращает список словарей с полями:
      ordinal, column, type, nullable, key, default, extra, comment
    """
    cursor.execute(
        "SELECT ORDINAL_POSITION, COLUMN_NAME, COLUMN_TYPE, IS_NULLABLE, "
        "       COLUMN_KEY, COLUMN_DEFAULT, EXTRA, COLUMN_COMMENT "
        "FROM information_schema.COLUMNS "
        "WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s "
        "ORDER BY ORDINAL_POSITION",
        (DB_NAME, table_name)
    )
    rows = cursor.fetchall()
    result = []
    for r in rows:
        result.append({
            'ordinal': r[0],
            'column': r[1],
            'type': r[2],
            'nullable': r[3],
            'key': r[4] or '',
            'default': '' if r[5] is None else str(r[5]),
            'extra': r[6] or '',
            'comment': r[7] or '',
        })
    return result


def get_table_comment(cursor, table_name):
    cursor.execute(
        "SELECT TABLE_COMMENT FROM information_schema.TABLES "
        "WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s",
        (DB_NAME, table_name)
    )
    row = cursor.fetchone()
    return row[0] if row and row[0] else ''


def get_row_count(cursor, table_name):
    cursor.execute(f"SELECT COUNT(*) FROM `{table_name}`")
    return cursor.fetchone()[0]


# ----------------------------------------------------------------------
# Выгрузка в Excel
# ----------------------------------------------------------------------
def export_to_excel(conn, tables_data, filename=OUT_XLSX):
    wb = openpyxl.Workbook()

    # --- Сводный лист ---
    ws_overview = wb.active
    ws_overview.title = "Обзор"

    # Стили
    title_font = Font(name='Arial', size=14, bold=True, color="FFFFFF")
    header_font = Font(name='Arial', size=11, bold=True, color="FFFFFF")
    normal_font = Font(name='Arial', size=11)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)

    title_fill = PatternFill("solid", fgColor="305496")
    header_fill = PatternFill("solid", fgColor="4472C4")
    alt_fill = PatternFill("solid", fgColor="D9E1F2")

    thin = Side(border_style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Заголовок
    ws_overview.merge_cells('A1:D1')
    ws_overview['A1'] = f"Структура базы данных «{DB_NAME}»"
    ws_overview['A1'].font = title_font
    ws_overview['A1'].fill = title_fill
    ws_overview['A1'].alignment = center
    ws_overview.row_dimensions[1].height = 30

    ws_overview.merge_cells('A2:D2')
    ws_overview['A2'] = f"Дата формирования: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws_overview['A2'].font = normal_font
    ws_overview['A2'].alignment = center

    # Заголовки сводной таблицы
    headers = ["Таблица", "Комментарий", "Кол-во полей", "Строк данных"]
    for col, h in enumerate(headers, start=1):
        cell = ws_overview.cell(row=4, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    row_idx = 5
    for tname, tdata in tables_data.items():
        ws_overview.cell(row=row_idx, column=1, value=tname).font = normal_font
        ws_overview.cell(row=row_idx, column=2,
                         value=tdata['comment'] or "—").font = normal_font
        ws_overview.cell(row=row_idx, column=3,
                         value=len(tdata['columns'])).font = normal_font
        ws_overview.cell(row=row_idx, column=4,
                         value=tdata['row_count']).font = normal_font

        for col in range(1, 5):
            c = ws_overview.cell(row=row_idx, column=col)
            c.border = border
            c.alignment = left if col in (1, 2) else center
            if row_idx % 2 == 1:
                c.fill = alt_fill
        row_idx += 1

    ws_overview.column_dimensions['A'].width = 20
    ws_overview.column_dimensions['B'].width = 45
    ws_overview.column_dimensions['C'].width = 15
    ws_overview.column_dimensions['D'].width = 15

    # --- Листы по таблицам ---
    for tname, tdata in tables_data.items():
        ws = wb.create_sheet(title=tname[:31])  # Excel ограничивает 31 символ

        ws.merge_cells('A1:H1')
        ws['A1'] = f"Таблица: {tname}"
        ws['A1'].font = title_font
        ws['A1'].fill = title_fill
        ws['A1'].alignment = center
        ws.row_dimensions[1].height = 25

        ws.merge_cells('A2:H2')
        ws['A2'] = tdata['comment'] or "—"
        ws['A2'].font = normal_font
        ws['A2'].alignment = center

        headers = ["№", "Поле", "Тип", "NULL", "Ключ",
                   "По умолчанию", "Доп.", "Комментарий"]
        for col, h in enumerate(headers, start=1):
            cell = ws.cell(row=4, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border

        row_idx = 5
        for col_data in tdata['columns']:
            ws.cell(row=row_idx, column=1, value=col_data['ordinal'])
            ws.cell(row=row_idx, column=2, value=col_data['column'])
            ws.cell(row=row_idx, column=3, value=col_data['type'])
            ws.cell(row=row_idx, column=4, value=col_data['nullable'])
            ws.cell(row=row_idx, column=5, value=col_data['key'])
            ws.cell(row=row_idx, column=6, value=col_data['default'])
            ws.cell(row=row_idx, column=7, value=col_data['extra'])
            ws.cell(row=row_idx, column=8, value=col_data['comment'])

            for col in range(1, 9):
                c = ws.cell(row=row_idx, column=col)
                c.font = normal_font
                c.border = border
                c.alignment = center if col in (1, 4, 5, 7) else left
                if row_idx % 2 == 1:
                    c.fill = alt_fill
            row_idx += 1

        widths = [6, 25, 25, 8, 8, 18, 12, 35]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(filename)
    print(f"[OK] Excel сохранён: {filename}")


# ----------------------------------------------------------------------
# Выгрузка в PDF
# ----------------------------------------------------------------------
def export_to_pdf(tables_data, filename=OUT_PDF):
    # Пытаемся зарегистрировать шрифт с поддержкой кириллицы
    cyrillic_font = "Helvetica"
    cyrillic_font_bold = "Helvetica-Bold"
    for fname, fpath in [
        ("DejaVu", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        ("DejaVu", "C:/Windows/Fonts/DejaVuSans.ttf"),
        ("Arial", "C:/Windows/Fonts/arial.ttf"),
    ]:
        if os.path.exists(fpath):
            try:
                pdfmetrics.registerFont(TTFont(fname, fpath))
                cyrillic_font = fname
                if fname == "Arial":
                    pdfmetrics.registerFont(TTFont("Arial-Bold", "C:/Windows/Fonts/arialbd.ttf"))
                    cyrillic_font_bold = "Arial-Bold"
                else:
                    cyrillic_font_bold = fname
                print(f"[i] Используется шрифт: {fpath}")
                break
            except Exception:
                continue

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'TitleRus', parent=styles['Title'],
        fontName=cyrillic_font_bold, fontSize=18, alignment=1,
        textColor=colors.HexColor('#305496'), spaceAfter=10
    )
    h2_style = ParagraphStyle(
        'H2Rus', parent=styles['Heading2'],
        fontName=cyrillic_font_bold, fontSize=14, alignment=0,
        textColor=colors.HexColor('#305496'), spaceAfter=6, spaceBefore=10
    )
    normal_style = ParagraphStyle(
        'NormalRus', parent=styles['Normal'],
        fontName=cyrillic_font, fontSize=9
    )

    doc = SimpleDocTemplate(
        filename, pagesize=landscape(A4),
        leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=15 * mm, bottomMargin=15 * mm
    )

    story = []
    story.append(Paragraph(f"Структура базы данных «{DB_NAME}»", title_style))
    story.append(Paragraph(
        f"Дата формирования: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        normal_style
    ))
    story.append(Spacer(1, 10))

    # --- Сводная таблица ---
    story.append(Paragraph("Обзор таблиц", h2_style))

    overview_data = [["Таблица", "Комментарий", "Полей", "Строк"]]
    for tname, tdata in tables_data.items():
        overview_data.append([
            tname,
            Paragraph(tdata['comment'] or "—", normal_style),
            str(len(tdata['columns'])),
            str(tdata['row_count'])
        ])

    overview_table = Table(overview_data,
                           colWidths=[60 * mm, 130 * mm, 25 * mm, 25 * mm],
                           repeatRows=1)
    overview_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4472C4')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), cyrillic_font_bold),
        ('FONTNAME', (0, 1), (-1, -1), cyrillic_font),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('ALIGN', (2, 1), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1),
         [colors.white, colors.HexColor('#D9E1F2')]),
    ]))
    story.append(overview_table)
    story.append(PageBreak())

    # --- Таблицы ---
    for tname, tdata in tables_data.items():
        story.append(Paragraph(f"Таблица: {tname}", h2_style))
        if tdata['comment']:
            story.append(Paragraph(f"Комментарий: {tdata['comment']}", normal_style))
        story.append(Spacer(1, 5))

        headers = ["№", "Поле", "Тип", "NULL", "Ключ",
                   "По умолчанию", "Доп.", "Комментарий"]
        data = [headers]
        for c in tdata['columns']:
            data.append([
                str(c['ordinal']),
                Paragraph(c['column'], normal_style),
                Paragraph(c['type'], normal_style),
                c['nullable'],
                c['key'],
                Paragraph(c['default'], normal_style),
                Paragraph(c['extra'], normal_style),
                Paragraph(c['comment'], normal_style),
            ])

        table = Table(
            data,
            colWidths=[10 * mm, 40 * mm, 45 * mm, 12 * mm, 12 * mm,
                       30 * mm, 25 * mm, 70 * mm],
            repeatRows=1
        )
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4472C4')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), cyrillic_font_bold),
            ('FONTNAME', (0, 1), (-1, -1), cyrillic_font),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
            ('ALIGN', (0, 1), (0, -1), 'CENTER'),
            ('ALIGN', (3, 1), (4, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('GRID', (0, 0), (-1, -1), 0.4, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1),
             [colors.white, colors.HexColor('#D9E1F2')]),
        ]))
        story.append(table)
        story.append(PageBreak())

    doc.build(story)
    print(f"[OK] PDF сохранён: {filename}")


# ----------------------------------------------------------------------
# Основная программа
# ----------------------------------------------------------------------
def main():
    print(f"=== Выгрузка структуры БД «{DB_NAME}» ===")
    print(f"Хост: {DB_HOST}, пользователь: {DB_USER}")

    try:
        conn = pymysql.connect(
            host=DB_HOST, user=DB_USER, password=DB_PASSWORD,
            database=DB_NAME, charset=DB_CHARSET, autocommit=True
        )
        cursor = conn.cursor()
        print("[OK] Подключение к БД установлено")
    except Exception as e:
        print(f"[ОШИБКА] Не удалось подключиться к БД: {e}")
        sys.exit(1)

    # Собираем структуру
    tables_data = {}
    for tname in TABLES:
        try:
            columns = get_table_structure(cursor, tname)
            if not columns:
                print(f"[!] Таблица {tname} не найдена в БД — пропущена")
                continue
            comment = get_table_comment(cursor, tname)
            row_count = get_row_count(cursor, tname)
            tables_data[tname] = {
                'columns': columns,
                'comment': comment,
                'row_count': row_count,
            }
            print(f"[i] {tname}: {len(columns)} полей, {row_count} строк")
        except Exception as e:
            print(f"[!] Ошибка при чтении {tname}: {e}")

    if not tables_data:
        print("[ОШИБКА] Не удалось прочитать ни одной таблицы")
        sys.exit(1)

    # Экспорт
    try:
        export_to_excel(conn, tables_data, OUT_XLSX)
    except Exception as e:
        print(f"[!] Ошибка Excel: {e}")

    try:
        export_to_pdf(tables_data, OUT_PDF)
    except Exception as e:
        print(f"[!] Ошибка PDF: {e}")

    conn.close()
    print("=== Готово ===")


if __name__ == "__main__":
    main()