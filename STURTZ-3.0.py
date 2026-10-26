import re
import os
import csv
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from openpyxl import load_workbook

# ============================================================
# STURTZ1.0-DEBUG v3
# ============================================================

BR_TO_PROFILE = {
    'BR-05000': '1032130010001', 'BR-01000': '1012130010001',
    'BR-03000': '1021030010001', 'BR-07000': '1012080010001',
    'BR-08000': '1022140010001', 'BR-09000': '1032320010001',
}

PROFILE_TYPES = {
    '1012130010001': 'рама', '1012080010001': 'рама', '1012330010001': 'рама',
    '1021030010001': 'импост', '1022140010001': 'импост', '1022350010001': 'импост',
    '1032130010001': 'створка', '1032320010001': 'створка', '1032350010001': 'створка',
}

WORK_NAMES = {
    '0101': 'Фреза 12 мм', '0201': 'Дренаж наружу вниз',
    '0301': 'Разметка T пластик', '0999': 'Фрез. под резину (рама/импост)',
    '1101': 'Дренаж фальца 45°', '1331': 'Разметка T металл',
    '1601': 'Дренаж фальца для створки', '1801': 'Дренаж наружу',
    '1803': 'Вентиляция фальца (створка)', '1804': 'Фрез. под резину (створка)',
    '1901': 'Фрез. 3 отв. под ручку', '4401': 'Шуруповёрт',
    '4501': 'Сверление под крепёж',
}

LOC_TO_STURTZ = {'T': 'O', 'B': 'U', 'H': 'W', 'V': 'S', 'L': 'L', 'R': 'R'}
LOCATION_MEANING = {
    'T': 'верх', 'B': 'низ', 'L': 'лево', 'R': 'право',
    'V': 'вертикальный', 'H': 'горизонтальный',
}

NEW_VER = 'Ver02.07-00'
SRC_VER = 'ZUH30'
S_CODE = '113292'
SRC_EXT = '.sob'
DST_EXT = '.sob'

NOTE_TRIGGER_DN = 'DN'
NOTE_TRIGGER_D25 = 'D25'

W4401_EDGE_BY_TYPE = {'рама': 1080, 'импост': 1080, 'створка': 1190}
W4401_EDGE_DEFAULT = 1080
W4401_MAX_STEP = 3000
W4401_MIN_SPAN = 1000
W4401_SHIFT = 200
W4401_TOL = 100
W4401_TAIL = '00000'

W1901_TAIL_DEFAULT = '00350'
W1901_TAIL_D25 = '00450'
W1901_HANDLE_OFFSET = 500

W0101_TAIL_FROM_0301 = '00000'
W0101_TAIL_FROM_1901 = '00600'

W1101_EDGE = 1500
W1101_OFFSET = 1000
W1101_L_LIMIT = 10000
W1101_TAIL = '00200'
W1101_IMPOST_L = 880
W1101_IMPOST_R = 820

W0201_EDGE = 1020
W0201_TAIL_STD = '00200'
W0201_TAIL_IMPOST = '00100'

W1601_EDGE = 1820
W1601_L_LIMIT = 10000
W1601_TAIL = '00200'

W1801_CENTER_OFFSET = 30
W1801_TAIL = '00200'

W_RUBBER_OFFSET = 380
W0999_TAIL = '00000'
W1804_TAIL = '00000'

W4501_BY_PROFILE = {
    '1012130010001': 550, '1012080010001': 550, '1012330010001': 577,
    '1032130010001': 720, '1032320010001': 730, '1032350010001': 700,
}
W4501_DEFAULT = 600
W4501_TAIL = '00000'
W4501_SKIP_TYPES = {'импост'}


# ============================================================
# ТРАНСЛИТЕРАЦИЯ
# ============================================================

TRANSLIT_TABLE = {
    'а':'a','б':'b','в':'v','г':'g','д':'d','е':'e','ё':'e','ж':'zh','з':'z',
    'и':'i','й':'j','к':'k','л':'l','м':'m','н':'n','о':'o','п':'p','р':'r',
    'с':'s','т':'t','у':'u','ф':'f','х':'h','ц':'c','ч':'ch','ш':'sh',
    'щ':'sch','ъ':'','ы':'y','ь':'','э':'e','ю':'ju','я':'ja',
    'А':'A','Б':'B','В':'V','Г':'G','Д':'D','Е':'E','Ё':'E','Ж':'Zh','З':'Z',
    'И':'I','Й':'J','К':'K','Л':'L','М':'M','Н':'N','О':'O','П':'P','Р':'R',
    'С':'S','Т':'T','У':'U','Ф':'F','Х':'H','Ц':'C','Ч':'Ch','Ш':'Sh',
    'Щ':'Sch','Ъ':'','Ы':'Y','Ь':'','Э':'E','Ю':'Ju','Я':'Ja',
}


def transliterate(text):
    if not text:
        return ''
    return ''.join(TRANSLIT_TABLE.get(ch, ch) for ch in text)


# ============================================================
# K-ПОЛЕ
# ============================================================

def normalize_order_from_xlsx(s):
    s = str(s or '').strip()
    if len(s) > 8:
        return s[4:12]
    return s.ljust(8)[:8]


def get_order_key(user_key, xlsx_order_normalized, source_key, src_filename):
    if user_key and user_key.strip():
        clean = re.sub(r'[^A-Za-z0-9]', '', user_key.strip())[:8]
        if clean:
            return clean.ljust(8)[:8]
    if xlsx_order_normalized:
        return xlsx_order_normalized.ljust(8)[:8]
    sk = (source_key or '').strip()
    if sk and sk != '00000000':
        return sk.ljust(8)[:8]
    if src_filename:
        base = os.path.splitext(os.path.basename(src_filename))[0]
        clean = re.sub(r'[^A-Za-z0-9]', '', base)[:8]
        if clean:
            return clean.ljust(8)[:8]
    return '00000000'


# ============================================================
# SOE
# ============================================================

def build_ednt(part_num, text):
    return f"EDNT{part_num:04d}T{text[:84]}"


def build_ednr(stab_num, text):
    return f"EDNR{stab_num:03d}T{text[:84]}"


def build_soe_for_part(part_num, rec):
    lines = []
    заказ = rec.get('заказ', '')
    конструкция = rec.get('конструкция', '')
    размер = rec.get('размер', '').replace('#', 'X')
    уплотнитель = rec.get('уплотнитель', '')
    ячейка = rec.get('ячейка', '')
    профиль = rec.get('профиль', '')
    цвет = rec.get('цвет', '')
    if цвет.startswith('Цвет:'):
        цвет = цвет[len('Цвет:'):]
    длина = rec.get('длина', '')
    артикул = rec.get('артикул', '')
    штапик = rec.get('штапик', '')
    if штапик.startswith('ШТАП:'):
        штапик = штапик[len('ШТАП:'):]
    фурнитура = rec.get('фурнитура', '')
    открывание = rec.get('открывание', '')
    примечание = rec.get('примечание', '')

    l1 = [transliterate(x) for x in [заказ, конструкция, размер, уплотнитель, ячейка] if x]
    if l1:
        lines.append(build_ednt(part_num, '#'.join(l1)))
    l2 = []
    if профиль: l2.append(transliterate(профиль))
    if цвет:    l2.append(transliterate(цвет))
    if длина:   l2.append(transliterate(длина))
    if артикул: l2.append(f"ST:{transliterate(артикул)}")
    if l2:
        lines.append(build_ednt(part_num, '#'.join(l2)))
    l3 = [transliterate(x) for x in [штапик, фурнитура, открывание] if x]
    if l3:
        lines.append(build_ednt(part_num, '#'.join(l3)))
    if примечание:
        lines.append(build_ednt(part_num, transliterate(примечание)))
    return lines


def build_soe_for_krl(stab_num, rec):
    lines = []
    krl = rec.get('krl', '')
    цвет = rec.get('цвет', '')
    if цвет.startswith('Цвет:'):
        цвет = цвет[len('Цвет:'):]
    if krl:
        lines.append(build_ednr(stab_num, transliterate(krl)))
    if цвет:
        lines.append(build_ednr(stab_num, transliterate(цвет)))
    return lines


# ============================================================
# XLSX
# ============================================================

def parse_xlsx_for_labels(path, log=print):
    wb = load_workbook(path, data_only=True)
    ws = wb['Лист1']
    flat = []
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            if cell is not None and str(cell).strip():
                flat.append(str(cell).strip())

    records = []
    current = None
    for s in flat:
        if re.match(r'^\d+-\d+-\d+$', s):
            if current:
                records.append(current)
            current = {'тип': 'деталь', 'заказ': s}
            continue
        if s.startswith('KRL:'):
            if current:
                records.append(current)
            color = ''
            for r in reversed(records):
                if r.get('тип') == 'деталь' and r.get('цвет'):
                    color = r['цвет']
                    break
            current = {'тип': 'остаток', 'krl': s, 'цвет': color}
            continue
        if current is None:
            continue
        if re.match(r'^K\d+\.\d+', s):
            current['конструкция'] = s
        elif s.startswith('ШТАП:'):
            current['штапик'] = s
        elif s.startswith('*R') and ('Штапик' in s or 'Shtapik' in s):
            current['штапик'] = s
        elif s.startswith('*R'):
            current['профиль'] = s
        elif 'VORNE' in s or s.startswith('*  '):
            current['фурнитура'] = s
        elif s.startswith('Упл:'):
            current['уплотнитель'] = s
        elif s.startswith('Цвет:'):
            current['цвет'] = s
        elif s.startswith('L:'):
            current['длина'] = s
        elif re.match(r'^\d+° \d+°$', s):
            current['углы'] = s
        elif re.match(r'^\d+\([а-яa-z]\)$', s):
            current['ячейка'] = s
        elif re.match(r'^\d+/[A-Za-zА-Яа-я0-9]+', s):
            current['артикул'] = s
        elif re.match(r'^\d+#\d+$', s):
            current['размер'] = s
        elif re.match(r'^[А-ЯA-Z]{1,3}\s*-\s*[А-Яа-яA-Za-z]{0,3}$', s):
            current['открывание'] = s
        else:
            current.setdefault('примечание', s)

    if current:
        records.append(current)
    log(f"XLSX: деталей = {sum(1 for r in records if r['тип']=='деталь')}, "
        f"остатков = {sum(1 for r in records if r['тип']=='остаток')}")
    return records


# ============================================================
# SOB СБОРЩИКИ
# ============================================================

def build_ksn(stab, length, profile_code, description, storage="000000"):
    return (f"KSN{stab:0>3}L{length:0>5}T{profile_code:<13}"
            f"B{description:<20}O{storage:<6}")


def build_ktn(part, cart, comp, order, pos, elem, location_stz,
              profile_code, description, length,
              g_val, d_val, s_val, w_val, i_val):
    return (f"KTN{part:0>4}C{cart:0>3}F{comp:0>3}K{order:<8}P{pos:<9}E{elem:0>3}"
            f"U{location_stz}T{profile_code:<13}B{description:<20}L{length:0>5}"
            f"G{g_val:<8}D{d_val}S{s_val}W{w_val}I{i_val:<18}S      Z")


def build_s(length):
    return f"S   /{S_CODE:<11}/{length:0>5}"


def compute_i(order, cart, comp, elem):
    try:
        comp_int = str(int(comp))
    except ValueError:
        comp_int = '0'
    return f"{order:<8}"[:8] + comp_int + f"{cart:0>3}" + "001" + f"{elem:0>3}"


# ============================================================
# SOB ПАРСЕРЫ
# ============================================================

def parse_ksn(line):
    m = re.match(r'^KSN(\d{3})L(\d{5})T\s+(BR-\d{5})(?:#\s*)?(.*?)\s*S?O(\d{6})', line)
    if m:
        pc = BR_TO_PROFILE.get(m.group(3))
        if not pc:
            return None
        return {'stab': m.group(1), 'length': m.group(2), 'profile_code': pc,
                'description': (m.group(4) or 'bel').strip().ljust(20)[:20],
                'storage': m.group(5)}
    m = re.match(r'^KSN(\d{3})L(\d{5})T(\S{13})B(.{20})O(\S{6})', line)
    if m:
        return {'stab': m.group(1), 'length': m.group(2),
                'profile_code': m.group(3), 'description': m.group(4),
                'storage': m.group(5)}
    return None


def parse_ktn(line):
    m = re.match(
        r'^KTN(\d{4})C(\d{3})F(\d{3})K(.{8})P(.{9})E(\d{3})'
        r'U([A-Z])T\s+(BR-\d{5})(?:#\s*)?(.*?)\s*S?L(\d{5})'
        r'G(\S{8})D(\S)S(\S)', line)
    if m:
        pc = BR_TO_PROFILE.get(m.group(8))
        if not pc:
            return None
        return {'part': m.group(1), 'cart': m.group(2), 'comp': m.group(3),
                'order': m.group(4).strip() or '00000000',
                'pos': m.group(5), 'elem': m.group(6),
                'location': m.group(7), 'profile_code': pc,
                'description': (m.group(9) or 'bel').strip().ljust(20)[:20],
                'length': m.group(10), 'g': m.group(11),
                'd': '1' if m.group(12) == '0' else m.group(12),
                's': m.group(13), 'w': '0', 'i': None}
    m = re.match(
        r'^KTN(\d{4})C(\d{3})F(\d{3})K(\S{8})P(\S{9})E(\d{3})'
        r'U([A-Z])T(\S{13})B(.{20})L(\d{5})G(\S{8})D(\S)S(\S)W(\S)I(\S{18})',
        line)
    if m:
        return {'part': m.group(1), 'cart': m.group(2), 'comp': m.group(3),
                'order': m.group(4), 'pos': m.group(5), 'elem': m.group(6),
                'location': m.group(7), 'profile_code': m.group(8),
                'description': m.group(9), 'length': m.group(10),
                'g': m.group(11), 'd': m.group(12), 's': m.group(13),
                'w': m.group(14), 'i': m.group(15)}
    return None


# ============================================================
# РАСЧЁТЫ
# ============================================================

def get_w4401_edge(pt):
    return W4401_EDGE_BY_TYPE.get(pt, W4401_EDGE_DEFAULT)


def compute_w4401_positions(L, edge):
    if L is None or L < 2 * edge:
        return []
    span = L - 2 * edge
    if span < W4401_MIN_SPAN:
        return [L // 2]
    N = max(1, (span + W4401_MAX_STEP - 1) // W4401_MAX_STEP)
    step = span / N
    return [int(round(edge + i * step)) for i in range(N + 1)]


def compute_w4401_with_handle(L, w1901_pos, edge):
    if L is None:
        return []
    left_end = w1901_pos - W1901_HANDLE_OFFSET
    right_beg = w1901_pos + W1901_HANDLE_OFFSET
    right_end = L - edge
    result = []
    if left_end > edge:
        span = left_end - edge
        N = max(1, (span + W4401_MAX_STEP - 1) // W4401_MAX_STEP)
        step = span / N
        for i in range(N + 1):
            result.append(int(round(edge + i * step)))
    else:
        result.append(edge)
    if right_end > right_beg:
        span = right_end - right_beg
        N = max(1, (span + W4401_MAX_STEP - 1) // W4401_MAX_STEP)
        step = span / N
        for i in range(N + 1):
            result.append(int(round(right_beg + i * step)))
    else:
        result.append(right_beg)
    return sorted(set(result))


def filter_w4401_handle(positions, w1901_pos):
    if w1901_pos is None:
        return positions
    zs = w1901_pos - W1901_HANDLE_OFFSET
    ze = w1901_pos + W1901_HANDLE_OFFSET
    return [p for p in positions if not (zs < p < ze)]


def apply_w0301_shift(positions, w0301_positions):
    if not w0301_positions:
        return positions
    out = []
    for p in positions:
        s = 0
        for q in w0301_positions:
            if abs(p - q) <= W4401_TOL:
                s = W4401_SHIFT
                break
        out.append(p + s)
    return out


def compute_w1101_positions(L, w0301):
    positions = [W1101_EDGE, L - W1101_EDGE]
    if w0301:
        for p in w0301:
            positions.append(p - W1101_OFFSET)
            positions.append(p + W1101_OFFSET)
    elif L > W1101_L_LIMIT:
        positions.append(L // 2)
    return sorted(set(p for p in positions if 0 < p < L))


def compute_w0201_positions(L):
    if L is None or L < 2 * W0201_EDGE:
        return []
    return [W0201_EDGE, L - W0201_EDGE]


def compute_w1601_positions(L):
    if L is None or L < 2 * W1601_EDGE:
        return []
    positions = [W1601_EDGE, L - W1601_EDGE]
    if L > W1601_L_LIMIT:
        positions.append(L // 2)
    return sorted(set(positions))


# ============================================================
# ОТЛАДОЧНЫЙ ОТЧЁТ
# ============================================================

def draw_scale(works, L, width=110):
    out = []
    step_chars = width // 10
    labels = ""
    for i in range(11):
        pos_mm = int(i * L / 10)
        label = str(pos_mm)
        if i < 10:
            label = label.ljust(step_chars)
        labels += label
    out.append(" " * 15 + labels)

    bar = ""
    for i in range(0, width + 1):
        if i % step_chars == 0:
            bar += "├" if i == 0 else ("┤" if i == width else "┼")
        else:
            bar += "─"
    out.append(" " * 15 + bar)

    for w in sorted(works, key=lambda x: min(x['positions_mm']) if x['positions_mm'] else 99999):
        line = list(bar)
        for pos in w['positions_mm']:
            col = round(pos / L * width) if L > 0 else 0
            col = max(0, min(width, col))
            line[col] = "▓"
        code_str = f"W{w['code']}"
        positions_str = ", ".join(f"{p:.0f}" for p in sorted(w['positions_mm']))
        out.append(f"{code_str:<7} " + "".join(line) + f"  [{positions_str}]")
    return out


def format_part_report(part):
    L = part['length_mm']
    out = []
    header = (f" KTN{part['ktn_num']:04d}  |  {part['profile_type']:<8} |  "
              f"L={L:8.1f} мм  |  LOC={part['location']} ({LOCATION_MEANING.get(part['location'], '?')})  "
              f"|  {part['profile_code']} ")
    out.append("┌" + "─" * 140 + "┐")
    out.append("│" + header.ljust(140) + "│")
    out.append("└" + "─" * 140 + "┘")
    out.append("")

    if part.get('sob_block'):
        out.append(f"  Стержень: KSN{part.get('ksn_stab', '???')}  "
                   f"(L={part.get('ksn_length_mm', 0):.0f} мм, "
                   f"{part.get('ksn_profile_code', '?')})")
        out.append("")
        out.append("  ─── Итоговый SOB-код для этой детали ───")
        for line in part['sob_block']:
            out.append(f"  {line}")
        out.append("")

    if not part['works']:
        out.append("  (нет работ)")
        return out

    out.append("  ┌──────┬───────┬─────────────────────────────────┬──────────────────────────────────────────────────┐")
    out.append("  │ Код  │ Хвост │ Название                        │ Позиции (мм)                                     │")
    out.append("  ├──────┼───────┼─────────────────────────────────┼──────────────────────────────────────────────────┤")
    for w in sorted(part['works'], key=lambda x: min(x['positions_mm']) if x['positions_mm'] else 99999):
        code = f"W{w['code']}"
        tail = w.get('tail', '')
        name = w['name'][:31]
        pos_str = ", ".join(f"{p:.0f}" for p in w['positions_mm'])
        if len(pos_str) > 48:
            pos_str = pos_str[:45] + "..."
        out.append(f"  │ {code:<4} │ {tail:<5} │ {name:<31} │ {pos_str:<48} │")
    out.append("  └──────┴───────┴─────────────────────────────────┴──────────────────────────────────────────────────┘")
    out.append("")

    if L > 0:
        out.append("  Визуализация расположения работ на детали:")
        out.append("")
        out.extend(draw_scale(part['works'], L, width=110))
        out.append("")
        out.append("  Легенда: ▓ = позиция работы, ├ ┤ = края детали, ┼ = метка шкалы")
        out.append("")

    return out


def format_debug_report(debug_parts, source_name=''):
    if not debug_parts:
        return "Нет данных."
    out = []
    total_works = sum(len(p['works']) for p in debug_parts)
    total_pos = sum(len(w['positions_mm']) for p in debug_parts for w in p['works'])
    out.append("╔" + "═" * 140 + "╗")
    out.append("║" + "ОТЛАДОЧНЫЙ ОТЧЁТ ПО РАБОТАМ  —  STURTZ1.0-DEBUG v3".center(140) + "║")
    out.append("║" + f"Файл: {source_name}".center(140) + "║")
    out.append("║" + f"Всего деталей: {len(debug_parts)}  |  Работ: {total_works}  |  Позиций: {total_pos}".center(140) + "║")
    out.append("╚" + "═" * 140 + "╝")
    out.append("")
    for part in debug_parts:
        out.extend(format_part_report(part))
        out.append("")
    return "\n".join(out)


# ============================================================
# ПРЕОБРАЗОВАНИЕ
# ============================================================

def transform_content(content, order_note='', user_order_key='',
                      src_filename='', add_w4501=False, log=print):
    note_DN = NOTE_TRIGGER_DN in (order_note or '')
    note_D25 = NOTE_TRIGGER_D25 in (order_note or '')
    log(f"Примечание: DN={note_DN}, D25={note_D25}")
    log(f"W1901 хвост: {'00450 (D25)' if note_D25 else '00350 (по умолчанию)'}")
    log(f"Работа 4501: {'вкл' if add_w4501 else 'выкл'}")

    lines = content.splitlines()
    out_sob = []
    debug_parts = []
    i = 0
    stab_counter = 0
    part_counter = 0
    current_ksn_stab = None
    current_ksn_length = None
    current_ksn_profile = None

    while i < len(lines):
        raw = lines[i].rstrip('\n').rstrip('\r')
        if not raw:
            i += 1
            continue
        if raw.strip() == SRC_VER:
            out_sob.append(NEW_VER)
            i += 1
            continue

        if raw.startswith('KSN'):
            ksn = parse_ksn(raw)
            stab_counter += 1
            if ksn:
                ksn['stab'] = f"{stab_counter:03d}"
                ksn_line = build_ksn(ksn['stab'], ksn['length'],
                                     ksn['profile_code'], ksn['description'],
                                     ksn['storage'])
                out_sob.append(ksn_line)
                current_ksn_stab = ksn['stab']
                current_ksn_length = int(ksn['length']) / 10.0 if ksn['length'] else 0
                current_ksn_profile = ksn['profile_code']
            i += 1
            continue

        if raw.startswith('KTN'):
            ktn = parse_ktn(raw)
            if not ktn:
                log(f"[!] Не распознан KTN: {raw[:80]}")
                i += 1
                continue
            profile_code = ktn['profile_code']
            profile_type = PROFILE_TYPES.get(profile_code)
            L = int(ktn['length']) if ktn['length'] else None
            location = ktn['location']
            part_counter += 1

            i += 1
            w_source = []
            while i < len(lines):
                nxt = lines[i].rstrip('\n').rstrip('\r')
                if not nxt:
                    i += 1
                    continue
                if nxt.startswith('W'):
                    w_source.append(nxt)
                    i += 1
                elif nxt.startswith('S'):
                    i += 1
                elif (nxt.startswith('KTN') or nxt.startswith('KSN')
                      or nxt.startswith('KRL') or nxt.startswith('Ver')):
                    break
                else:
                    i += 1

            if L:
                w_source = [w for w in w_source
                            if not (lambda m: m and int(m.group(1)) >= L)
                            (re.match(r'W\d+/00/(\d+)/', w))]

            w0301_positions = []
            w1901_position = None
            existing = {k: set() for k in
                        ('0999', '1101', '0201', '1601', '1801', '1804')}
            for w in w_source:
                mw = re.match(r'W0301/00/(\d+)/', w)
                if mw: w0301_positions.append(int(mw.group(1)))
                mw = re.match(r'W1901/00/(\d+)/', w)
                if mw and w1901_position is None:
                    w1901_position = int(mw.group(1))
                for code in existing:
                    mw = re.match(rf'W{code}/00/(\d+)/', w)
                    if mw: existing[code].add(int(mw.group(1)))

            order_key_final = get_order_key(user_order_key, '', ktn['order'], src_filename)
            ktn['order'] = order_key_final
            ktn['i'] = compute_i(order_key_final, ktn['cart'], ktn['comp'], ktn['elem'])

            # ====================== W-элементы ======================
            w_items = []
            for w in w_source:
                # W1901 — ВСЕГДА пересобираем с правильным хвостом
                m1901_mod = re.match(r'W1901/00/(\d+)/', w)
                if m1901_mod:
                    pos = int(m1901_mod.group(1))
                    tail = W1901_TAIL_D25 if note_D25 else W1901_TAIL_DEFAULT
                    w = f"W1901/00/{pos:05d}/{tail}"

                mw = re.match(r'W(\d+)/00/(\d+)/', w)
                if mw:
                    w_items.append((int(mw.group(2)), mw.group(1), w))

                m0301 = re.match(r'W0301/00/(\d+)/', w)
                if m0301:
                    pos = int(m0301.group(1))
                    w_items.append((pos, '0101',
                                    f"W0101/00/{pos:05d}/{W0101_TAIL_FROM_0301}"))
                m1901 = re.match(r'W1901/00/(\d+)/', w)
                if m1901:
                    pos = int(m1901.group(1))
                    w_items.append((pos, '0101',
                                    f"W0101/00/{pos:05d}/{W0101_TAIL_FROM_1901}"))

            if add_w4501 and L and profile_type not in W4501_SKIP_TYPES:
                pos = W4501_BY_PROFILE.get(profile_code, W4501_DEFAULT)
                if pos < L:
                    w_items.append((pos, '4501',
                                    f"W4501/00/{pos:05d}/{W4501_TAIL}"))

            if L:
                if profile_type == 'импост' and location == 'H':
                    for p in (W1101_IMPOST_L, L - W1101_IMPOST_R):
                        if 0 < p < L and p not in existing['1101']:
                            w_items.append((p, '1101',
                                            f"W1101/00/{p:05d}/{W1101_TAIL}"))
                            existing['1101'].add(p)
                elif profile_type == 'рама' and location == 'B':
                    for p in compute_w1101_positions(L, w0301_positions):
                        if 0 < p < L and p not in existing['1101']:
                            w_items.append((p, '1101',
                                            f"W1101/00/{p:05d}/{W1101_TAIL}"))
                            existing['1101'].add(p)

            add_w0201 = ((profile_type in ('створка', 'рама') and location in ('T', 'B'))
                         or (profile_type == 'импост' and location == 'V'))
            frame_B_DN = (profile_type == 'рама' and location == 'B' and note_DN)
            if L and add_w0201:
                if frame_B_DN:
                    for p in compute_w0201_positions(L):
                        if p not in existing['1801']:
                            w_items.append((p, '1801',
                                            f"W1801/00/{p:05d}/{W1801_TAIL}"))
                            existing['1801'].add(p)
                else:
                    tail = W0201_TAIL_IMPOST if profile_type == 'импост' else W0201_TAIL_STD
                    for p in compute_w0201_positions(L):
                        if p not in existing['0201']:
                            w_items.append((p, '0201',
                                            f"W0201/00/{p:05d}/{tail}"))
                            existing['0201'].add(p)

            if L and ((profile_type == 'створка' and location in ('T', 'B'))
                      or (profile_type == 'рама' and location == 'B')):
                for p in compute_w1601_positions(L):
                    if p not in existing['1601']:
                        w_items.append((p, '1601',
                                        f"W1601/00/{p:05d}/{W1601_TAIL}"))
                        existing['1601'].add(p)

            if L and profile_type == 'импост' and location == 'H':
                p = L // 2 + W1801_CENTER_OFFSET
                if 0 < p < L and p not in existing['1801']:
                    w_items.append((p, '1801',
                                    f"W1801/00/{p:05d}/{W1801_TAIL}"))
                    existing['1801'].add(p)

            if L and w0301_positions:
                for p in w0301_positions:
                    for delta in (-W_RUBBER_OFFSET, W_RUBBER_OFFSET):
                        new_p = p + delta
                        if not (0 < new_p < L):
                            continue
                        if profile_type in ('рама', 'импост'):
                            if new_p not in existing['0999']:
                                w_items.append((new_p, '0999',
                                                f"W0999/00/{new_p:05d}/{W0999_TAIL}"))
                                existing['0999'].add(new_p)
                        elif profile_type == 'створка':
                            if new_p not in existing['1804']:
                                w_items.append((new_p, '1804',
                                                f"W1804/00/{new_p:05d}/{W1804_TAIL}"))
                                existing['1804'].add(new_p)

            edge = get_w4401_edge(profile_type)
            if L:
                if w1901_position is not None:
                    positions = compute_w4401_with_handle(L, w1901_position, edge)
                else:
                    positions = compute_w4401_positions(L, edge)
                positions = filter_w4401_handle(positions, w1901_position)
                positions = apply_w0301_shift(positions, w0301_positions)
                for p in positions:
                    if 0 < p < L:
                        w_items.append((p, '4401',
                                        f"W4401/00/{p:05d}/{W4401_TAIL}"))

            # Отладка: собираем данные с хвостом
            works_by_code = {}
            for pos10, code, wline in w_items:
                m = re.match(rf'W{code}/00/\d+/(\d+)', wline)
                tail = m.group(1) if m else ''
                if code not in works_by_code:
                    works_by_code[code] = {'tail': tail, 'positions': []}
                works_by_code[code]['positions'].append(pos10 / 10.0)

            # SOB-блок
            loc_stz = LOC_TO_STURTZ.get(location, 'L')
            ktn_out = build_ktn(
                part_counter, ktn['cart'], ktn['comp'],
                ktn['order'], ktn['pos'], ktn['elem'], loc_stz,
                profile_code, ktn['description'], ktn['length'],
                ktn['g'], ktn['d'], ktn['s'], ktn['w'], ktn['i'])
            s_out = build_s(ktn['length'])
            w_items_sorted = sorted(w_items, key=lambda x: (x[0], x[1]))
            w_lines_out = [w_line for _, _, w_line in w_items_sorted]
            sob_block = [ktn_out] + w_lines_out + [s_out]

            part_debug = {
                'ktn_num': part_counter,
                'profile_code': profile_code,
                'profile_type': profile_type or '?',
                'location': location or '?',
                'length_mm': L / 10.0 if L else 0,
                'works': [
                    {'code': code,
                     'name': WORK_NAMES.get(code, '?'),
                     'tail': info['tail'],
                     'positions_mm': sorted(info['positions'])}
                    for code, info in works_by_code.items()
                ],
                'sob_block': sob_block,
                'ksn_stab': current_ksn_stab,
                'ksn_length_mm': current_ksn_length,
                'ksn_profile_code': current_ksn_profile,
            }
            debug_parts.append(part_debug)

            log(f"── KTN{part_counter:04d} | {profile_type} | L={L/10 if L else 0:.0f} | LOC={location}")
            for w in sorted(part_debug['works'], key=lambda x: x['code']):
                pos_str = ", ".join(f"{p:.0f}" for p in w['positions_mm'])
                log(f"     W{w['code']} tail={w['tail']} {w['name']:<30} [{pos_str}]")

            out_sob.append(ktn_out)
            for w_line in w_lines_out:
                out_sob.append(w_line)
            out_sob.append(s_out)
            continue

        if raw.startswith('KRL'):
            out_sob.append(raw)
            i += 1
            continue
        if raw.startswith('S'):
            i += 1
            continue
        out_sob.append(raw)
        i += 1

    sob_result = '\n'.join(out_sob) + '\n'
    return sob_result, '', debug_parts


# ============================================================
# GUI
# ============================================================

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("STURTZ1.0-DEBUG v3 — отладка работ")
        self.geometry("1000x760")
        self.resizable(True, True)

        self.src_var = tk.StringVar()
        self.dst_var = tk.StringVar()
        self.note_var = tk.StringVar()
        self.order_key_var = tk.StringVar()
        self.add_4501_var = tk.BooleanVar(value=False)
        self.location_var = tk.StringVar(value="Готово к запуску")
        self.last_debug_parts = []
        self.last_src_name = ''

        pad = {'padx': 8, 'pady': 4}

        tk.Label(self, text="Файл-источник (.sob):").grid(row=0, column=0, sticky='w', **pad)
        tk.Entry(self, textvariable=self.src_var, width=74).grid(row=0, column=1, sticky='we', **pad)
        tk.Button(self, text="Обзор...", command=self.choose_src).grid(row=0, column=2, **pad)

        tk.Label(self, text="Результирующий .sob:").grid(row=1, column=0, sticky='w', **pad)
        tk.Entry(self, textvariable=self.dst_var, width=74).grid(row=1, column=1, sticky='we', **pad)
        tk.Button(self, text="Обзор...", command=self.choose_dst).grid(row=1, column=2, **pad)

        tk.Label(self, text="Код заказа (K, до 8 знаков):").grid(row=2, column=0, sticky='w', **pad)
        tk.Entry(self, textvariable=self.order_key_var, width=74).grid(row=2, column=1, sticky='we', **pad)

        tk.Label(self, text="Примечание заказа:").grid(row=3, column=0, sticky='w', **pad)
        tk.Entry(self, textvariable=self.note_var, width=74).grid(row=3, column=1, sticky='we', **pad)

        tk.Label(self, text="(триггеры: DN → W1801 на раме B; D25 → W1901 хвост 00450)",
                 fg="#666", font=("Arial", 8)).grid(row=4, column=1, sticky='w', **pad)

        tk.Checkbutton(self, text="Добавить работу 4501",
                       variable=self.add_4501_var).grid(row=5, column=1, sticky='w', **pad)

        tk.Label(self, textvariable=self.location_var,
                 font=("Arial", 11, "bold"), fg="#1565C0").grid(row=6, column=1, sticky='w', **pad)

        btn_frame = tk.Frame(self)
        btn_frame.grid(row=7, column=1, columnspan=2, sticky='e', **pad)

        tk.Button(btn_frame, text="Сгенерировать (SOB + SOE)",
                  command=self.run_generation,
                  bg="#2E7D32", fg="white",
                  font=("Arial", 10, "bold")).pack(side='left', padx=4)
        tk.Button(btn_frame, text="Показать работы на деталях",
                  command=self.show_works,
                  bg="#FF6F00", fg="white",
                  font=("Arial", 10, "bold")).pack(side='left', padx=4)
        tk.Button(btn_frame, text="Экспорт отчёта в .txt",
                  command=self.export_report).pack(side='left', padx=4)

        tk.Label(self, text="Журнал:").grid(row=8, column=0, sticky='nw', **pad)
        self.log_text = tk.Text(self, height=22, wrap='word')
        self.log_text.grid(row=8, column=1, columnspan=2, sticky='nsew', **pad)
        scroll = ttk.Scrollbar(self, command=self.log_text.yview)
        scroll.grid(row=8, column=3, sticky='ns')
        self.log_text.config(yscrollcommand=scroll.set)

        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(8, weight=1)

    def log(self, msg):
        self.log_text.insert('end', msg + '\n')
        self.log_text.see('end')
        self.update_idletasks()

    def choose_src(self):
        path = filedialog.askopenfilename(
            title="Выберите .sob файл",
            filetypes=[("Файлы центра (.sob)", "*.sob"), ("Все файлы", "*.*")])
        if not path:
            return
        self.src_var.set(path)
        folder = os.path.dirname(path)
        base = os.path.splitext(os.path.basename(path))[0]
        self.dst_var.set(os.path.join(folder, f"{base}_out{DST_EXT}"))

    def choose_dst(self):
        path = filedialog.asksaveasfilename(
            title="Сохранить .sob как",
            defaultextension=DST_EXT,
            filetypes=[("Файлы центра (.sob)", "*.sob"), ("Все файлы", "*.*")],
            initialfile=os.path.basename(self.dst_var.get()) or f"output{DST_EXT}",
            initialdir=os.path.dirname(self.dst_var.get()) or None)
        if path:
            if not path.lower().endswith(DST_EXT):
                path += DST_EXT
            self.dst_var.set(path)

    def run_generation(self):
        self.log_text.delete('1.0', 'end')
        self.location_var.set("Обработка...")
        src = self.src_var.get().strip()
        dst = self.dst_var.get().strip()
        note = self.note_var.get().strip()
        okey = self.order_key_var.get().strip()
        add_4501 = bool(self.add_4501_var.get())

        if not src or not os.path.isfile(src):
            messagebox.showerror("Ошибка", "Файл-источник .sob не найден.")
            return
        base = os.path.splitext(src)[0]
        xlsx_path = base + '.xlsx'
        if not os.path.isfile(xlsx_path):
            xlsx_path = filedialog.askopenfilename(
                title=f"Не найден {os.path.basename(base)}.xlsx — выберите вручную",
                filetypes=[("Excel", "*.xlsx"), ("Все файлы", "*.*")],
                initialdir=os.path.dirname(src))
            if not xlsx_path:
                return
        if not dst:
            messagebox.showerror("Ошибка", "Не указан путь для результата.")
            return
        if not dst.lower().endswith(DST_EXT):
            dst += DST_EXT
            self.dst_var.set(dst)
        dst_soe = os.path.splitext(dst)[0] + '.soe'

        try:
            self.log(f"SOB:  {src}")
            self.log(f"XLSX: {xlsx_path}")
            self.log("")
            with open(src, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()

            xlsx_records = parse_xlsx_for_labels(xlsx_path, log=self.log)

            self.log("")
            sob_result, _, debug_parts = transform_content(
                content, order_note=note, user_order_key=okey,
                src_filename=src, add_w4501=add_4501, log=self.log)

            self.last_debug_parts = debug_parts
            self.last_src_name = os.path.basename(src)

            with open(dst, 'w', encoding='utf-8', newline='\r\n') as f:
                f.write(sob_result)

            xlsx_dets = [r for r in xlsx_records if r.get('тип') == 'деталь']
            xlsx_krls = [r for r in xlsx_records if r.get('тип') == 'остаток']
            soe_lines = []
            for idx, rec in enumerate(xlsx_dets, 1):
                soe_lines.extend(build_soe_for_part(idx, rec))
            for idx, rec in enumerate(xlsx_krls, 1):
                soe_lines.extend(build_soe_for_krl(idx, rec))
            soe_result = '\r\n'.join(soe_lines) + '\r\n' if soe_lines else ''
            with open(dst_soe, 'w', encoding='utf-8', newline='\r\n') as f:
                f.write(soe_result)

            self.location_var.set(
                f"Готово: деталей={len(debug_parts)}, "
                f"работ={sum(len(p['works']) for p in debug_parts)}")
            self.log("")
            self.log("═══ Готово ═══")
            self.log(f"SOB: {dst}")
            self.log(f"SOE: {dst_soe}")
            self.log(f"Деталей: {len(debug_parts)}")
            messagebox.showinfo("Успех",
                f"Файлы сохранены:\n{dst}\n{dst_soe}\n\n"
                f"Деталей: {len(debug_parts)}\n"
                f"Нажмите «Показать работы на деталях» для отчёта.")
        except Exception as e:
            self.log(f"[ОШИБКА] {e}")
            import traceback
            self.log(traceback.format_exc())
            messagebox.showerror("Ошибка", str(e))

    def show_works(self):
        if not self.last_debug_parts:
            messagebox.showwarning("Нет данных", "Сначала сгенерируйте выгрузку.")
            return
        win = tk.Toplevel(self)
        win.title("Отладка — работы на деталях + SOB-код")
        win.geometry("1500x900")

        text = tk.Text(win, wrap='none', font=('Consolas', 9))
        text.grid(row=0, column=0, sticky='nsew')
        vsb = ttk.Scrollbar(win, orient='vertical', command=text.yview)
        hsb = ttk.Scrollbar(win, orient='horizontal', command=text.xview)
        text.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        vsb.grid(row=0, column=1, sticky='ns')
        hsb.grid(row=1, column=0, sticky='ew')

        report = format_debug_report(self.last_debug_parts, self.last_src_name)
        text.insert('1.0', report)
        text.config(state='disabled')

        win.grid_rowconfigure(0, weight=1)
        win.grid_columnconfigure(0, weight=1)

        btn = tk.Frame(win)
        btn.grid(row=2, column=0, columnspan=2, pady=6)
        tk.Button(btn, text="Экспорт отчёта в .txt",
                  command=self.export_report).pack(side='left', padx=4)
        tk.Button(btn, text="Копировать всё в буфер",
                  command=lambda: self._copy_to_clipboard(report)).pack(side='left', padx=4)
        tk.Button(btn, text="Закрыть", command=win.destroy).pack(side='left', padx=4)

    def _copy_to_clipboard(self, text):
        self.clipboard_clear()
        self.clipboard_append(text)
        messagebox.showinfo("Скопировано", "Отчёт скопирован в буфер обмена.")

    def export_report(self):
        if not self.last_debug_parts:
            messagebox.showwarning("Нет данных", "Сначала запустите генерацию.")
            return
        path = filedialog.asksaveasfilename(
            title="Сохранить отчёт",
            defaultextension=".txt",
            filetypes=[("Text", "*.txt"), ("Все файлы", "*.*")],
            initialfile="debug_report.txt")
        if not path:
            return
        report = format_debug_report(self.last_debug_parts, self.last_src_name)
        with open(path, 'w', encoding='utf-8') as f:
            f.write(report)
        self.log(f"Отчёт сохранён: {path}")
        messagebox.showinfo("Готово", f"Отчёт сохранён:\n{path}")


if __name__ == '__main__':
    App().mainloop()