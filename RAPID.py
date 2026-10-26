import re
import os
import json
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from openpyxl import load_workbook

# ============================================================
# RAPID1.0 — генератор выгрузки для MFC Evolution
# ============================================================

# --- Профили ST → PRGNR RAPID ---
PROFILE_TO_PRGN = {
    '1012080010001': '717000',   # 70 РАМА
    '1012130010001': '716000',   # 60 РАМА
    '1012330010001': '716000',   # 60 РАМА (альт.)
    '1021030010001': '736000',   # 60 ИМПОСТ
    '1022140010001': '737000',   # 70 ИМПОСТ
    '1022350010001': '737000',   # 70 ИМПОСТ (белый)
    '1032130010001': '726000',   # 60 Створка
    '1032320010001': '727000',   # 70 Створка
    '1032350010001': '726000',   # 60 Створка (белый)
}

BR_TO_PROFILE = {
    'BR-05000': '1032130010001',
    'BR-01000': '1012130010001',
    'BR-03000': '1021030010001',
    'BR-07000': '1012080010001',
    'BR-08000': '1022140010001',
    'BR-09000': '1032320010001',
}

PROFILE_TYPES = {
    '1012130010001': 'рама', '1012080010001': 'рама', '1012330010001': 'рама',
    '1021030010001': 'импост', '1022140010001': 'импост', '1022350010001': 'импост',
    '1032130010001': 'створка', '1032320010001': 'створка', '1032350010001': 'створка',
}

# --- Работы ST → RAPID opcodes ---
W_COMMON_MAP = {
    '0201': '7000',   # дренаж наружу вниз
    '0301': '7030',   # разметка T пластик
    '1331': '7018',   # разметка T металл
    '1801': '7002',   # дренаж наружу
    '4401': '7020',   # шуруповёрт
    '4501': '7020',   # сверление под крепёж (в тот же блок)
}

W_BY_TYPE = {
    'рама': {
        '0101': '7181',
        '1101': '7004',
        '0999': '7014',
        '0903': '7016',
    },
    'импост': {
        '0101': '7181',
        '1101': '7004',
        '0999': '7014',
        '0903': '7016',
    },
    'створка': {
        '0101': '7182',
        '1601': '7004',
        '1804': '7015',
        '1803': '7017',
    },
}

# --- W1901 → RAPID opcodes ---
W1901_DEFAULT = ['7160', '7180']
W1901_D25     = ['7165', '7180']
W1901_2SIDE   = ['7165', '7166', '7180']

NOTE_DN     = 'DN'
NOTE_D25    = 'D25'
NOTE_25D    = '25D'
NOTE_2SIDE  = '2х стор ручка'

# --- Штрих-код ---
BARCODE_START    = '0000900901'
BARCODE_LOG_FILE = 'rapid_barcode.json'

# --- Постоянные для L-полей ---
L5_DEFAULT   = '00540'    # глубина вставки армировки
L6_SUBTRACT  = 1400       # L6 = L1 − 1400 (в 0.1 мм)

# --- Геометрия операций ---
W4401_EDGE_BY_TYPE = {'рама': 1080, 'импост': 1080, 'створка': 1190}
W4401_EDGE_DEFAULT = 1080
W4401_MAX_STEP     = 3000
W4401_MIN_SPAN     = 1000
W4401_SHIFT        = 200
W4401_TOL          = 100
W1901_HANDLE_OFFSET = 500

W0201_EDGE = 1020
W1601_EDGE = 1820
W1601_L_LIMIT = 10000

W1101_EDGE = 1500
W1101_OFFSET = 1000
W1101_L_LIMIT = 10000
W1101_IMPOST_L = 880
W1101_IMPOST_R = 820

W_RUBBER_OFFSET = 380
W1801_CENTER_OFFSET = 30

W4501_BY_PROFILE = {
    '1012130010001': 550, '1012080010001': 550, '1012330010001': 577,
    '1032130010001': 720, '1032320010001': 730, '1032350010001': 700,
}
W4501_DEFAULT = 600
W4501_SKIP_TYPES = {'импост'}


# ============================================================
# ШТРИХ-КОД
# ============================================================

class BarcodeManager:
    def __init__(self, log_path):
        self.log_path = log_path
        self.value = self._load()

    def _load(self):
        if os.path.isfile(self.log_path):
            try:
                with open(self.log_path, 'r', encoding='utf-8') as f:
                    data = json.load(f)
                    v = str(data.get('last', BARCODE_START))
                    if len(v) == 10 and v.isdigit():
                        return v
            except Exception:
                pass
        return BARCODE_START

    def next(self):
        current = self.value
        try:
            self.value = f"{int(self.value) + 1:010d}"
        except ValueError:
            self.value = BARCODE_START
        return current

    def save(self):
        with open(self.log_path, 'w', encoding='utf-8') as f:
            json.dump({'last': self.value}, f, ensure_ascii=False, indent=2)

    def reset(self):
        self.value = BARCODE_START


# ============================================================
# ПАРСЕР XLSX
# ============================================================

def parse_xlsx(path, log=print):
    wb = load_workbook(path, data_only=True)
    ws = wb['Лист1']
    flat = []
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            if cell is not None and str(cell).strip():
                flat.append(str(cell).strip())

    records = []
    current = None
    for s in flat:
        if re.match(r'^\d+-\d+-\d+$', s):
            if current:
                records.append(current)
            current = {'тип': 'деталь', 'заказ': s}
            continue
        if s.startswith('KRL:'):
            if current:
                records.append(current)
            current = {'тип': 'остаток', 'krl': s}
            continue
        if current is None:
            continue

        if re.match(r'^K\d+\.\d+', s):
            current['конструкция'] = s
        elif s.startswith('ШТАП:'):
            current['штапик'] = s
        elif s.startswith('*R') and ('Штапик' in s or 'Shtapik' in s):
            current['штапик'] = s
        elif s.startswith('*R'):
            current['профиль'] = s
        elif 'VORNE' in s or s.startswith('*  '):
            current['фурнитура'] = s
        elif s.startswith('Упл:'):
            current['уплотнитель'] = s
        elif s.startswith('Цвет:'):
            current['цвет'] = s
        elif s.startswith('L:'):
            current['длина'] = s
        elif re.match(r'^\d+° \d+°$', s):
            current['углы'] = s
        elif re.match(r'^\d+\([а-яa-z]\)$', s):
            current['ячейка'] = s
        elif re.match(r'^\d+/[A-Za-zА-Яа-я0-9]+', s):
            current['артикул'] = s
        elif re.match(r'^\d+#\d+$', s):
            current['размер'] = s
        elif re.match(r'^[А-ЯA-Z]{1,3}\s*-\s*[А-Яа-яA-Za-z]{0,3}$', s):
            current['открывание'] = s
        else:
            current.setdefault('примечание', s)

    if current:
        records.append(current)

    log(f"XLSX: деталей = {sum(1 for r in records if r['тип']=='деталь')}, "
        f"остатков = {sum(1 for r in records if r['тип']=='остаток')}")
    return records


# ============================================================
# УГЛЫ И STEEL
# ============================================================

def parse_angles(s):
    if not s:
        return (45, 45)
    m = re.match(r'^(\d+)°\s+(\d+)°$', s.strip())
    if m:
        return (int(m.group(1)), int(m.group(2)))
    return (45, 45)


def angle_code(a1, a2):
    return {
        (45, 45): '00',
        (90, 45): '01',
        (45, 90): '02',
        (90, 90): '03',
    }.get((a1, a2), '00')


def steel_code(articolo):
    if not articolo:
        return '00000'
    digits = re.sub(r'\D', '', articolo)
    if not digits:
        return '00000'
    return digits[:5].ljust(5, '0')


# ============================================================
# ПАРСЕРЫ SOB
# ============================================================

def parse_ksn(line):
    m = re.match(r'^KSN(\d{3})L(\d{5})T\s+(BR-\d{5})(?:#\s*)?(.*?)\s*S?O(\d{6})', line)
    if m:
        pc = BR_TO_PROFILE.get(m.group(3))
        if not pc:
            return None
        return {'stab': m.group(1), 'length': m.group(2), 'profile_code': pc,
                'description': (m.group(4) or 'bel').strip().ljust(20)[:20]}
    m = re.match(r'^KSN(\d{3})L(\d{5})T(\S{13})B(.{20})O(\S{6})', line)
    if m:
        return {'stab': m.group(1), 'length': m.group(2),
                'profile_code': m.group(3), 'description': m.group(4)}
    return None


def parse_ktn(line):
    m = re.match(
        r'^KTN(\d{4})C(\d{3})F(\d{3})K(.{8})P(.{9})E(\d{3})'
        r'U([A-Z])T\s+(BR-\d{5})(?:#\s*)?(.*?)\s*S?L(\d{5})'
        r'G(\S{8})D(\S)S(\S)', line)
    if m:
        pc = BR_TO_PROFILE.get(m.group(8))
        if not pc:
            return None
        return {'part': m.group(1), 'cart': m.group(2), 'comp': m.group(3),
                'order': m.group(4).strip() or '00000000',
                'pos': m.group(5), 'elem': m.group(6), 'location': m.group(7),
                'profile_code': pc,
                'description': (m.group(9) or 'bel').strip().ljust(20)[:20],
                'length': m.group(10), 'g': m.group(11),
                'd': '1' if m.group(12) == '0' else m.group(12), 's': m.group(13)}
    m = re.match(
        r'^KTN(\d{4})C(\d{3})F(\d{3})K(\S{8})P(\S{9})E(\d{3})'
        r'U([A-Z])T(\S{13})B(.{20})L(\d{5})G(\S{8})D(\S)S(\S)W(\S)I(\S{18})',
        line)
    if m:
        return {'part': m.group(1), 'cart': m.group(2), 'comp': m.group(3),
                'order': m.group(4), 'pos': m.group(5), 'elem': m.group(6),
                'location': m.group(7), 'profile_code': m.group(8),
                'description': m.group(9), 'length': m.group(10),
                'g': m.group(11), 'd': m.group(12), 's': m.group(13)}
    return None


# ============================================================
# РАСЧЁТЫ ПОЗИЦИЙ
# ============================================================

def get_w4401_edge(pt):
    return W4401_EDGE_BY_TYPE.get(pt, W4401_EDGE_DEFAULT)


def compute_w4401_positions(L, edge):
    if L is None or L < 2 * edge:
        return []
    span = L - 2 * edge
    if span < W4401_MIN_SPAN:
        return [L // 2]
    N = max(1, (span + W4401_MAX_STEP - 1) // W4401_MAX_STEP)
    step = span / N
    return [int(round(edge + i * step)) for i in range(N + 1)]


def compute_w4401_with_handle(L, w1901_pos, edge):
    if L is None:
        return []
    left_end  = w1901_pos - W1901_HANDLE_OFFSET
    right_beg = w1901_pos + W1901_HANDLE_OFFSET
    right_end = L - edge
    result = []
    if left_end > edge:
        span = left_end - edge
        N = max(1, (span + W4401_MAX_STEP - 1) // W4401_MAX_STEP)
        step = span / N
        for i in range(N + 1):
            result.append(int(round(edge + i * step)))
    else:
        result.append(edge)
    if right_end > right_beg:
        span = right_end - right_beg
        N = max(1, (span + W4401_MAX_STEP - 1) // W4401_MAX_STEP)
        step = span / N
        for i in range(N + 1):
            result.append(int(round(right_beg + i * step)))
    else:
        result.append(right_beg)
    return sorted(set(result))


def filter_w4401_handle(positions, w1901_pos):
    if w1901_pos is None:
        return positions
    zs = w1901_pos - W1901_HANDLE_OFFSET
    ze = w1901_pos + W1901_HANDLE_OFFSET
    return [p for p in positions if not (zs < p < ze)]


def apply_w0301_shift(positions, w0301_positions):
    if not w0301_positions:
        return positions
    out = []
    for p in positions:
        s = 0
        for q in w0301_positions:
            if abs(p - q) <= W4401_TOL:
                s = W4401_SHIFT
                break
        out.append(p + s)
    return out


def compute_w1101_positions(L, w0301):
    positions = [W1101_EDGE, L - W1101_EDGE]
    if w0301:
        for p in w0301:
            positions.append(p - W1101_OFFSET)
            positions.append(p + W1101_OFFSET)
    elif L > W1101_L_LIMIT:
        positions.append(L // 2)
    return sorted(set(p for p in positions if 0 < p < L))


def compute_w0201_positions(L):
    if L is None or L < 2 * W0201_EDGE:
        return []
    return [W0201_EDGE, L - W0201_EDGE]


def compute_w1601_positions(L):
    if L is None or L < 2 * W1601_EDGE:
        return []
    positions = [W1601_EDGE, L - W1601_EDGE]
    if L > W1601_L_LIMIT:
        positions.append(L // 2)
    return sorted(set(positions))


# ============================================================
# СБОРЩИКИ RAPID-СТРОК
# ============================================================

def build_p_row(job, batch, prgnr, count=1):
    job_num = re.sub(r'\D', '', str(job))[-6:].zfill(6)
    return f"P{job_num}{batch:0>4}{str(prgnr):<10}{count:0>3}"


def build_n_row(packet, barcode, cut_code, l_fields):
    head = f"N{packet:0>3}C{str(barcode):<10}Z001A01{cut_code}0"
    return head + ''.join(l_fields)


def build_l_op(opcode, positions):
    if not positions:
        return ''
    n = min(len(positions), 9)   # максимум 9 (1 знак count)
    s = f"L{opcode}{n:1d}"
    for p in sorted(set(positions))[:n]:
        s += f"L{p:0>5}"
    return s


# ============================================================
# ГЛАВНЫЙ КОНВЕРТЕР
# ============================================================

def transform_to_rapid(content, xlsx_records,
                       order_note='', user_order_key='',
                       add_w4501=False, barcode_mgr=None, log=print):
    note_DN    = NOTE_DN     in (order_note or '')
    note_D25   = NOTE_D25    in (order_note or '')
    note_25D   = NOTE_25D    in (order_note or '')
    note_2side = NOTE_2SIDE  in (order_note or '')

    log(f"Примечания: DN={note_DN}, D25={note_D25}, "
        f"25D={note_25D}, 2-side={note_2side}")

    xlsx_dets = [r for r in xlsx_records if r.get('тип') == 'деталь']

    lines = content.splitlines()
    parts = []
    i = 0
    det_idx = 0

    while i < len(lines):
        raw = lines[i].rstrip('\n').rstrip('\r')
        if not raw:
            i += 1
            continue
        if raw.strip() == 'ZUH30' or raw.startswith('KSN') or raw.startswith('KRL'):
            i += 1
            continue

        if not raw.startswith('KTN'):
            i += 1
            continue

        ktn = parse_ktn(raw)
        if not ktn:
            log(f"[!] Не распознан KTN: {raw[:60]}")
            i += 1
            continue

        profile_code = ktn['profile_code']
        profile_type = PROFILE_TYPES.get(profile_code)
        L = int(ktn['length']) if ktn['length'] else None
        location = ktn['location']

        # Сбор W-блока
        i += 1
        w_source = []
        while i < len(lines):
            nxt = lines[i].rstrip('\n').rstrip('\r')
            if not nxt:
                i += 1
                continue
            if nxt.startswith('W'):
                w_source.append(nxt)
                i += 1
            elif nxt.startswith('S'):
                i += 1
            elif (nxt.startswith('KTN') or nxt.startswith('KSN')
                  or nxt.startswith('KRL') or nxt.startswith('ZUH')):
                break
            else:
                i += 1

        if L:
            w_source = [w for w in w_source
                        if not (lambda m: m and int(m.group(1)) >= L)
                        (re.match(r'W\d+/00/(\d+)/', w))]

        # Сбор данных
        w0301_positions = []
        w1901_position = None
        for w in w_source:
            mw = re.match(r'W0301/00/(\d+)/', w)
            if mw: w0301_positions.append(int(mw.group(1)))
            mw = re.match(r'W1901/00/(\d+)/', w)
            if mw and w1901_position is None:
                w1901_position = int(mw.group(1))

        # {opcode: [positions]}
        ops = {}

        def add_op(code, pos):
            ops.setdefault(code, []).append(pos)

        # W0101 от W0301 / W1901
        for w in w_source:
            m0301 = re.match(r'W0301/00/(\d+)/', w)
            if m0301:
                add_op(W_BY_TYPE[profile_type]['0101'], int(m0301.group(1)))
            m1901 = re.match(r'W1901/00/(\d+)/', w)
            if m1901:
                add_op(W_BY_TYPE[profile_type]['0101'], int(m1901.group(1)))

        # W0201 / W1801
        add_w0201 = (
            (profile_type in ('створка', 'рама') and location in ('T', 'B')) or
            (profile_type == 'импост' and location == 'V')
        )
        frame_B_DN = (profile_type == 'рама' and location == 'B' and note_DN)
        if L and add_w0201:
            for p in compute_w0201_positions(L):
                if frame_B_DN:
                    add_op('7002', p)
                else:
                    add_op(W_COMMON_MAP['0201'], p)

        # W1101
        if L:
            if profile_type == 'импост' and location == 'H':
                for p in (W1101_IMPOST_L, L - W1101_IMPOST_R):
                    if 0 < p < L:
                        add_op(W_BY_TYPE[profile_type]['1101'], p)
            elif profile_type == 'рама' and location == 'B':
                for p in compute_w1101_positions(L, w0301_positions):
                    add_op(W_BY_TYPE[profile_type]['1101'], p)

        # W1601 (створка)
        if L and profile_type == 'створка' and location in ('T', 'B'):
            for p in compute_w1601_positions(L):
                add_op(W_BY_TYPE['створка']['1601'], p)

        # W1901 → RAPID opcodes
        if w1901_position is not None:
            if note_2side:
                opcodes = W1901_2SIDE
            elif note_25D:
                opcodes = W1901_D25
            else:
                opcodes = W1901_DEFAULT
            for opc in opcodes:
                add_op(opc, w1901_position)

        # W0999 / W1804 от W0301 ± 380
        if L and w0301_positions:
            for p in w0301_positions:
                for delta in (-W_RUBBER_OFFSET, W_RUBBER_OFFSET):
                    new_p = p + delta
                    if not (0 < new_p < L):
                        continue
                    if profile_type in ('рама', 'импост'):
                        add_op(W_BY_TYPE[profile_type]['0999'], new_p)
                    elif profile_type == 'створка':
                        add_op(W_BY_TYPE['створка']['1804'], new_p)

        # W4401 + W4501 → оба 7020
        edge = get_w4401_edge(profile_type)
        if L:
            if w1901_position is not None:
                positions = compute_w4401_with_handle(L, w1901_position, edge)
            else:
                positions = compute_w4401_positions(L, edge)
            positions = filter_w4401_handle(positions, w1901_position)
            positions = apply_w0301_shift(positions, w0301_positions)
            for p in positions:
                if 0 < p < L:
                    add_op('7020', p)

        if add_w4501 and L and profile_type not in W4501_SKIP_TYPES:
            pos = W4501_BY_PROFILE.get(profile_code, W4501_DEFAULT)
            if pos < L:
                add_op('7020', pos)

        # XLSX-данные
        xls = xlsx_dets[det_idx] if det_idx < len(xlsx_dets) else {}
        det_idx += 1

        a1, a2 = parse_angles(xls.get('углы', ''))
        cut_code = angle_code(a1, a2)
        l4 = steel_code(xls.get('артикул', ''))

        parts.append({
            'profile_code': profile_code,
            'profile_type': profile_type,
            'prgnr': PROFILE_TO_PRGN.get(profile_code, '000000'),
            'length': L,
            'angles': (a1, a2),
            'cut_code': cut_code,
            'l4': l4,
            'ops': ops,
            'location': location,
        })

    # Группировка по PRGNR
    parts.sort(key=lambda p: (p['prgnr'], p.get('location', '')))

    out = []
    batch = 0
    job_num = re.sub(r'\D', '', str(user_order_key or ''))[-6:].zfill(6)

    for prgnr in sorted(set(p['prgnr'] for p in parts)):
        batch += 1
        group = [p for p in parts if p['prgnr'] == prgnr]

        out.append(build_p_row(job_num, batch, prgnr, 1))

        for pkt, part in enumerate(group, 1):
            barcode = barcode_mgr.next() if barcode_mgr else '0000000000'

            L1 = part['length']
            L4 = part['l4']
            L6 = max(0, L1 - L6_SUBTRACT)
            a1, a2 = part['angles']
            L11 = f"{a1 * 100:05d}"
            L12 = f"{a2 * 100:05d}"

            l_fields = [
                f"L{L1:0>5}",
                "L00000",            # L2
                "L00000",            # L3
                f"L{L4:0>5}",        # L4
                f"L{L5_DEFAULT}",    # L5
                f"L{L6:0>5}",        # L6
                "L00000",            # L7
                "L00000",            # L8
                "L00000",            # L9
                "L00000",            # L10
                f"L{L11}",           # L11
                f"L{L12}",           # L12
                "L00000",            # L13
                "L00000",            # L14
            ]

            for opcode in sorted(part['ops'].keys()):
                l_fields.append(build_l_op(opcode, part['ops'][opcode]))

            out.append(build_n_row(pkt, barcode, part['cut_code'], l_fields))

    if barcode_mgr:
        barcode_mgr.save()

    return '\n'.join(out) + '\n'


# ============================================================
# GUI
# ============================================================

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("RAPID1.0 — MFC Evolution")
        self.geometry("960x740")
        self.resizable(True, True)

        self.src_var = tk.StringVar()
        self.dst_var = tk.StringVar()
        self.note_var = tk.StringVar()
        self.order_key_var = tk.StringVar()
        self.add_4501_var = tk.BooleanVar(value=False)
        self.barcode_var = tk.StringVar(value="—")

        self.barcode_mgr = BarcodeManager(BARCODE_LOG_FILE)

        pad = {'padx': 8, 'pady': 4}

        tk.Label(self, text="Файл-источник (.sob):").grid(row=0, column=0, sticky='w', **pad)
        tk.Entry(self, textvariable=self.src_var, width=74).grid(row=0, column=1, sticky='we', **pad)
        tk.Button(self, text="Обзор...", command=self.choose_src).grid(row=0, column=2, **pad)

        tk.Label(self, text="Результирующий .ubz:").grid(row=1, column=0, sticky='w', **pad)
        tk.Entry(self, textvariable=self.dst_var, width=74).grid(row=1, column=1, sticky='we', **pad)
        tk.Button(self, text="Обзор...", command=self.choose_dst).grid(row=1, column=2, **pad)

        tk.Label(self, text="Номер заказа:").grid(row=2, column=0, sticky='w', **pad)
        tk.Entry(self, textvariable=self.order_key_var, width=74).grid(row=2, column=1, sticky='we', **pad)
        tk.Label(self, text="(6 цифр, войдёт в P-строку)", fg="#666",
                 font=("Arial", 8)).grid(row=3, column=1, sticky='w', **pad)

        tk.Label(self, text="Примечание заказа:").grid(row=4, column=0, sticky='w', **pad)
        tk.Entry(self, textvariable=self.note_var, width=74).grid(row=4, column=1, sticky='we', **pad)
        tk.Label(self, text="(триггеры: DN, D25, 25D, '2х стор ручка')",
                 fg="#666", font=("Arial", 8)).grid(row=5, column=1, sticky='w', **pad)

        tk.Checkbutton(self, text="Добавить работу 4501",
                       variable=self.add_4501_var).grid(row=6, column=1, sticky='w', **pad)

        tk.Label(self, text="Следующий штрих-код:").grid(row=7, column=0, sticky='w', **pad)
        tk.Label(self, textvariable=self.barcode_var,
                 font=("Arial", 11, "bold"), fg="#1565C0").grid(row=7, column=1, sticky='w', **pad)

        btn_frame = tk.Frame(self)
        btn_frame.grid(row=8, column=1, columnspan=2, sticky='e', **pad)

        tk.Button(btn_frame, text="Сгенерировать .ubz",
                  command=self.run_generation,
                  bg="#2E7D32", fg="white",
                  font=("Arial", 10, "bold")).pack(side='left', padx=4)
        tk.Button(btn_frame, text="Сбросить штрих-код",
                  command=self.reset_barcode).pack(side='left', padx=4)

        tk.Label(self, text="Журнал:").grid(row=9, column=0, sticky='nw', **pad)
        self.log_text = tk.Text(self, height=22, wrap='word')
        self.log_text.grid(row=9, column=1, columnspan=2, sticky='nsew', **pad)
        scroll = ttk.Scrollbar(self, command=self.log_text.yview)
        scroll.grid(row=9, column=3, sticky='ns')
        self.log_text.config(yscrollcommand=scroll.set)

        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(9, weight=1)

        self.update_barcode_label()

    def log(self, msg):
        self.log_text.insert('end', msg + '\n')
        self.log_text.see('end')
        self.update_idletasks()

    def update_barcode_label(self):
        self.barcode_var.set(self.barcode_mgr.value)

    def reset_barcode(self):
        self.barcode_mgr.reset()
        self.barcode_mgr.save()
        self.update_barcode_label()
        self.log(f"Штрих-код сброшен на {BARCODE_START}")

    def choose_src(self):
        path = filedialog.askopenfilename(
            title="Выберите .sob файл",
            filetypes=[("Файлы центра (.sob)", "*.sob"), ("Все файлы", "*.*")])
        if not path:
            return
        self.src_var.set(path)
        folder = os.path.dirname(path)
        base = os.path.splitext(os.path.basename(path))[0]
        self.dst_var.set(os.path.join(folder, f"{base}.ubz"))

    def choose_dst(self):
        path = filedialog.asksaveasfilename(
            title="Сохранить .ubz как",
            defaultextension=".ubz",
            filetypes=[("RAPID files", "*.ubz"), ("Все файлы", "*.*")],
            initialfile=os.path.basename(self.dst_var.get()) or "output.ubz",
            initialdir=os.path.dirname(self.dst_var.get()) or None)
        if path:
            self.dst_var.set(path)

    def run_generation(self):
        self.log_text.delete('1.0', 'end')
        src = self.src_var.get().strip()
        dst = self.dst_var.get().strip()
        note = self.note_var.get().strip()
        okey = self.order_key_var.get().strip()
        add_4501 = bool(self.add_4501_var.get())

        if not src or not os.path.isfile(src):
            messagebox.showerror("Ошибка", "Файл-источник .sob не найден.")
            return
        base = os.path.splitext(src)[0]
        xlsx_path = base + '.xlsx'
        if not os.path.isfile(xlsx_path):
            xlsx_path = filedialog.askopenfilename(
                title=f"Не найден {os.path.basename(base)}.xlsx — выберите вручную",
                filetypes=[("Excel", "*.xlsx"), ("Все файлы", "*.*")],
                initialdir=os.path.dirname(src))
            if not xlsx_path:
                return

        if not dst:
            messagebox.showerror("Ошибка", "Не указан путь для результата.")
            return

        try:
            self.log(f"SOB:  {src}")
            self.log(f"XLSX: {xlsx_path}")
            self.log(f"UBZ:  {dst}")

            with open(src, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()

            xlsx_records = parse_xlsx(xlsx_path, log=self.log)

            result = transform_to_rapid(
                content, xlsx_records,
                order_note=note, user_order_key=okey,
                add_w4501=add_4501,
                barcode_mgr=self.barcode_mgr,
                log=self.log)

            with open(dst, 'w', encoding='utf-8', newline='\r\n') as f:
                f.write(result)

            self.update_barcode_label()
            self.log("Готово.")
            messagebox.showinfo("Успех", f"Файл сохранён:\n{dst}")
        except Exception as e:
            self.log(f"[ОШИБКА] {e}")
            import traceback
            self.log(traceback.format_exc())
            messagebox.showerror("Ошибка", str(e))


if __name__ == '__main__':
    App().mainloop()